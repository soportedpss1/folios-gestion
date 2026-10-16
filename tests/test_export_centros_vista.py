"""Exportaciones del detalle de centro divididas por pestaña (?vista=...)."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models.devolucion_folio import DevolucionFolio
from app.models.entrega_folio import EntregaFolio
from tests.conftest import make_folio

Y = 2026

VISTAS = ('pendientes', 'todos', 'devueltos')


def _seed(tipo, centro, operador):
    """Un folio entregado pendiente (7001) y uno devuelto (7002)."""
    pendiente = make_folio(tipo, anio=Y, numero=7001, estado='entregado')
    db.session.add(EntregaFolio(folio_id=pendiente.id, fechaEntrega=date(Y, 6, 1),
                                centroId=centro.id, userId=operador.id))
    devuelto = make_folio(tipo, anio=Y, numero=7002, estado='devuelto')
    db.session.add(EntregaFolio(folio_id=devuelto.id, fechaEntrega=date(Y, 6, 2),
                                centroId=centro.id, userId=operador.id))
    db.session.add(DevolucionFolio(folio_id=devuelto.id,
                                   fechaDevolucion=date(Y, 6, 10),
                                   centroId=centro.id, userId=operador.id))
    db.session.commit()
    return pendiente, devuelto


def _nombres_hojas(resp):
    wb = load_workbook(BytesIO(resp.data))
    return wb, wb.sheetnames


# --- Excel ---


def test_excel_vista_pendientes_una_sola_hoja(auth_client, app, tipo, centro, operador):
    _seed(tipo, centro, operador)
    resp = auth_client.get(f'/centros/{centro.id}/export/excel?vista=pendientes')
    assert resp.status_code == 200
    wb, hojas = _nombres_hojas(resp)
    assert hojas == ['Pendientes'], f'esperaba solo la hoja Pendientes, hay {hojas}'
    filas = list(wb['Pendientes'].iter_rows(values_only=True))
    numeros = [f[0] for f in filas[1:]]
    assert 7001 in numeros, 'el pendiente debe exportarse'
    assert 7002 not in numeros, 'el devuelto no debe estar en Pendientes'


def test_excel_vista_todos_una_sola_hoja(auth_client, app, tipo, centro, operador):
    _seed(tipo, centro, operador)
    resp = auth_client.get(f'/centros/{centro.id}/export/excel?vista=todos')
    assert resp.status_code == 200
    wb, hojas = _nombres_hojas(resp)
    assert hojas == ['Todos Entregados'], f'esperaba solo "Todos Entregados", hay {hojas}'
    filas = list(wb['Todos Entregados'].iter_rows(values_only=True))
    numeros = [f[0] for f in filas[1:]]
    assert 7001 in numeros and 7002 in numeros


def test_excel_vista_devueltos_una_sola_hoja(auth_client, app, tipo, centro, operador):
    _seed(tipo, centro, operador)
    resp = auth_client.get(f'/centros/{centro.id}/export/excel?vista=devueltos')
    assert resp.status_code == 200
    wb, hojas = _nombres_hojas(resp)
    assert hojas == ['Devueltos'], f'esperaba solo la hoja Devueltos, hay {hojas}'
    filas = list(wb['Devueltos'].iter_rows(values_only=True))
    numeros = [f[0] for f in filas[1:]]
    assert 7002 in numeros
    assert 7001 not in numeros, 'el pendiente no debe estar en Devueltos'


def test_excel_sin_vista_sigue_exportando_las_tres_hojas(
    auth_client, app, tipo, centro, operador
):
    """Regresión: sin parámetro, el export completo no cambia."""
    _seed(tipo, centro, operador)
    resp = auth_client.get(f'/centros/{centro.id}/export/excel')
    assert resp.status_code == 200
    _, hojas = _nombres_hojas(resp)
    assert hojas == ['Pendientes', 'Todos Entregados', 'Devueltos']


def test_excel_vista_invalida_es_400(auth_client, centro):
    resp = auth_client.get(f'/centros/{centro.id}/export/excel?vista=otra')
    assert resp.status_code == 400


# --- PDF ---


def _fake_html(monkeypatch):
    import weasyprint

    class _FakeHTML:
        captured = []

        def __init__(self, string=None, **kwargs):
            _FakeHTML.captured.append(string)

        def write_pdf(self):
            return b'%PDF-1.4 fake'

    _FakeHTML.captured = []
    monkeypatch.setattr(weasyprint, 'HTML', _FakeHTML)
    return _FakeHTML


def _weasyprint_disponible():
    try:
        import weasyprint  # noqa: F401
    except OSError as exc:
        pytest.skip(f'weasyprint sin libpango (solo Docker): {exc}')


def test_pdf_vista_pendientes_solo_su_seccion(
    auth_client, app, tipo, centro, operador, monkeypatch
):
    _weasyprint_disponible()
    _seed(tipo, centro, operador)
    fake = _fake_html(monkeypatch)

    resp = auth_client.get(f'/centros/{centro.id}/export/pdf?vista=pendientes')
    assert resp.status_code == 200
    assert len(fake.captured) == 1
    html = fake.captured[0]
    assert 'Pendientes de Devolución' in html
    assert '<td>7001</td>' in html, 'fila del pendiente ausente'
    assert 'Todos Entregados (' not in html, 'no debe incluir la sección Todos'
    assert 'Devueltos (' not in html, 'no debe incluir la sección Devueltos'


def test_pdf_vista_devueltos_solo_su_seccion(
    auth_client, app, tipo, centro, operador, monkeypatch
):
    _weasyprint_disponible()
    _seed(tipo, centro, operador)
    fake = _fake_html(monkeypatch)

    resp = auth_client.get(f'/centros/{centro.id}/export/pdf?vista=devueltos')
    assert resp.status_code == 200
    html = fake.captured[0]
    assert 'Devueltos (' in html
    assert '<td>7002</td>' in html, 'fila del devuelto ausente'
    assert 'Pendientes de Devolución' not in html
    assert 'Todos Entregados (' not in html


def test_pdf_sin_vista_incluye_las_tres_secciones(
    auth_client, app, tipo, centro, operador, monkeypatch
):
    """Regresión: sin parámetro, el PDF completo no cambia."""
    _weasyprint_disponible()
    _seed(tipo, centro, operador)
    fake = _fake_html(monkeypatch)

    resp = auth_client.get(f'/centros/{centro.id}/export/pdf')
    assert resp.status_code == 200
    html = fake.captured[0]
    assert 'Pendientes de Devolución' in html
    assert 'Todos Entregados (' in html
    assert 'Devueltos (' in html


def test_pdf_vista_invalida_es_400(auth_client, centro):
    """Validación ANTES de importar weasyprint: 400 sin requerir pango."""
    resp = auth_client.get(f'/centros/{centro.id}/export/pdf?vista=otra')
    assert resp.status_code == 400


# --- template: los botones siguen la pestaña activa ---


def test_detalle_botones_exportan_pestaña_activa(auth_client, app, centro):
    resp = auth_client.get(f'/centros/{centro.id}/detail')
    assert resp.status_code == 200
    body = resp.get_data(as_text=True)
    assert 'id="btn-export-excel"' in body, 'botón Excel sin id para el JS'
    assert 'id="btn-export-pdf"' in body, 'botón PDF sin id para el JS'
    assert '/export/excel?vista=pendientes' in body, \
        'href inicial debe apuntar a la pestaña activa (pendientes)'
    assert '/export/pdf?vista=pendientes' in body
    assert 'shown.bs.tab' in body, 'falta el JS que cambia los href al cambiar de pestaña'
