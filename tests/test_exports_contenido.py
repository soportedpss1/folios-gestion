"""Fase 4 — guardas de contenido de exportaciones (refactors eager-loading/join)."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models.entrega_folio import EntregaFolio
from app.models.recepcion_folio import RecepcionFolio
from tests.conftest import make_folio

Y = 2026


def _seed_recepcion(tipo, operador, rango_id=1, ini=100, fin=110):
    r = RecepcionFolio(fecha=date(Y, 3, 5), anioCert=Y, tipoCert_id=tipo.id,
                       folioInicial=ini, folioFinal=fin, rangoId=rango_id,
                       userId=operador.id)
    db.session.add(r)
    db.session.commit()
    return r


def test_reportes_excel_contiene_filas_correctas(auth_client, app, tipo, centro, operador):
    folio = make_folio(tipo, anio=Y, numero=6001, estado='entregado', digitado=True)
    db.session.add(EntregaFolio(folio_id=folio.id, fechaEntrega=date(Y, 5, 10),
                                centroId=centro.id, userId=operador.id))
    db.session.commit()

    resp = auth_client.get(f'/reportes/export/excel?anioCert={Y}')
    assert resp.status_code == 200

    wb = load_workbook(BytesIO(resp.data))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == 'Folio'
    assert len(rows) == 2, f'esperaba header + 1 fila, hay {len(rows)}'
    assert rows[1][0] == 6001
    assert rows[1][1] == tipo.name
    assert rows[1][3] == centro.name
    assert rows[1][4] == 'entregado'
    assert rows[1][5] == 'Sí'
    assert rows[1][8] == '10/05/2026'


def test_centros_excel_hoja_pendientes_correcta(auth_client, app, tipo, centro, operador):
    folio = make_folio(tipo, anio=Y, numero=6002, estado='entregado')
    db.session.add(EntregaFolio(folio_id=folio.id, fechaEntrega=date(Y, 5, 11),
                                centroId=centro.id, userId=operador.id))
    db.session.commit()

    resp = auth_client.get(f'/centros/{centro.id}/export/excel')
    assert resp.status_code == 200

    wb = load_workbook(BytesIO(resp.data))
    ws = wb['Pendientes']
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == 'Folio'
    assert len(rows) == 2, f'esperaba header + 1 fila, hay {len(rows)}'
    assert rows[1][0] == 6002
    assert rows[1][3] == operador.username


def test_reportes_pdf_contiene_filas(auth_client, app, tipo, centro, operador, monkeypatch):
    try:
        import weasyprint
    except OSError as exc:
        pytest.skip(f'weasyprint sin libpango (solo Docker): {exc}')

    class _FakeHTML:
        captured = []

        def __init__(self, string=None, **kwargs):
            _FakeHTML.captured.append(string)

        def write_pdf(self):
            return b'%PDF-1.4 fake'

    folio = make_folio(tipo, anio=Y, numero=6003, estado='entregado')
    db.session.add(EntregaFolio(folio_id=folio.id, fechaEntrega=date(Y, 5, 12),
                                centroId=centro.id, userId=operador.id))
    db.session.commit()
    _FakeHTML.captured = []
    monkeypatch.setattr(weasyprint, 'HTML', _FakeHTML)

    resp = auth_client.get(f'/reportes/export/pdf?anioCert={Y}')
    assert resp.status_code == 200
    assert len(_FakeHTML.captured) == 1
    html = _FakeHTML.captured[0]
    assert '<td>6003</td>' in html, 'fila del folio ausente en PDF'
    assert centro.name in html


def test_recepcion_index_muestra_botones_export(auth_client):
    resp = auth_client.get('/recepcion/')
    assert resp.status_code == 200
    assert b'/recepcion/export/excel' in resp.data, 'botón Excel ausente'
    assert b'/recepcion/export/pdf' in resp.data, 'botón PDF ausente'


def test_recepcion_excel_contiene_filas(auth_client, app, tipo, operador):
    r = _seed_recepcion(tipo, operador)

    resp = auth_client.get('/recepcion/export/excel')
    assert resp.status_code == 200
    assert 'recepciones.xlsx' in resp.headers['Content-Disposition']

    wb = load_workbook(BytesIO(resp.data))
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0] == ('#', 'Fecha', 'Año Cert.', 'Tipo', 'Folio Inicial',
                       'Folio Final', 'Total Folios', 'Registrado Por')
    assert len(rows) == 2, f'esperaba header + 1 fila, hay {len(rows)}'
    assert rows[1][0] == r.id
    assert rows[1][1] == '05/03/2026'
    assert rows[1][2] == Y
    assert rows[1][3] == tipo.name
    assert rows[1][4] == 100
    assert rows[1][5] == 110
    assert rows[1][6] == 11
    assert rows[1][7] == operador.username


def test_recepcion_pdf_contiene_filas(auth_client, app, tipo, operador, monkeypatch):
    try:
        import weasyprint
    except OSError as exc:
        pytest.skip(f'weasyprint sin libpango (solo Docker): {exc}')

    class _FakeHTML:
        captured = []

        def __init__(self, string=None, **kwargs):
            _FakeHTML.captured.append(string)

        def write_pdf(self):
            return b'%PDF-1.4 fake'

    _seed_recepcion(tipo, operador, ini=200, fin=205)
    _FakeHTML.captured = []
    monkeypatch.setattr(weasyprint, 'HTML', _FakeHTML)

    resp = auth_client.get('/recepcion/export/pdf')
    assert resp.status_code == 200
    assert len(_FakeHTML.captured) == 1
    html = _FakeHTML.captured[0]
    assert '<td>200</td>' in html, 'folio inicial ausente en PDF'
    assert '<td>205</td>' in html, 'folio final ausente en PDF'
    assert tipo.name in html
    assert operador.username in html


def test_recepcion_pdf_escapa_html(auth_client, app, tipo, operador, monkeypatch):
    try:
        import weasyprint
    except OSError as exc:
        pytest.skip(f'weasyprint sin libpango (solo Docker): {exc}')

    class _FakeHTML:
        captured = []

        def __init__(self, string=None, **kwargs):
            _FakeHTML.captured.append(string)

        def write_pdf(self):
            return b'%PDF-1.4 fake'

    tipo.name = '<script>alert(1)</script>'
    db.session.commit()
    _seed_recepcion(tipo, operador, ini=300, fin=301)
    _FakeHTML.captured = []
    monkeypatch.setattr(weasyprint, 'HTML', _FakeHTML)

    resp = auth_client.get('/recepcion/export/pdf')
    assert resp.status_code == 200
    html = _FakeHTML.captured[0]
    assert '<script>' not in html, 'HTML sin escapar en PDF'
    assert '&lt;script&gt;' in html, 'escape ausente en PDF'


def test_centros_index_muestra_botones_export(auth_client):
    resp = auth_client.get('/centros/')
    assert resp.status_code == 200
    assert b'/centros/export/excel' in resp.data, 'botón Excel ausente'
    assert b'/centros/export/pdf' in resp.data, 'botón PDF ausente'


def test_centros_lista_excel_contiene_filas(auth_client, app, centro):
    resp = auth_client.get('/centros/export/excel')
    assert resp.status_code == 200
    assert 'centros.xlsx' in resp.headers['Content-Disposition']

    wb = load_workbook(BytesIO(resp.data))
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0] == ('#', 'Nombre', 'Teléfono', 'Dirección', 'Estado')
    assert len(rows) == 2, f'esperaba header + 1 fila, hay {len(rows)}'
    assert rows[1][0] == centro.id
    assert rows[1][1] == centro.name
    assert rows[1][2] == centro.telefono
    assert rows[1][3] == centro.direccion
    assert rows[1][4] == 'Activo'


def test_centros_lista_pdf_contiene_filas(auth_client, app, centro, monkeypatch):
    try:
        import weasyprint
    except OSError as exc:
        pytest.skip(f'weasyprint sin libpango (solo Docker): {exc}')

    class _FakeHTML:
        captured = []

        def __init__(self, string=None, **kwargs):
            _FakeHTML.captured.append(string)

        def write_pdf(self):
            return b'%PDF-1.4 fake'

    _FakeHTML.captured = []
    monkeypatch.setattr(weasyprint, 'HTML', _FakeHTML)

    resp = auth_client.get('/centros/export/pdf')
    assert resp.status_code == 200
    assert 'centros.pdf' in resp.headers['Content-Disposition']
    assert len(_FakeHTML.captured) == 1
    html = _FakeHTML.captured[0]
    assert centro.name in html
    assert centro.direccion in html


def test_centros_lista_pdf_escapa_html(auth_client, app, centro, monkeypatch):
    try:
        import weasyprint
    except OSError as exc:
        pytest.skip(f'weasyprint sin libpango (solo Docker): {exc}')

    class _FakeHTML:
        captured = []

        def __init__(self, string=None, **kwargs):
            _FakeHTML.captured.append(string)

        def write_pdf(self):
            return b'%PDF-1.4 fake'

    centro.name = '<script>alert(1)</script>'
    db.session.commit()
    _FakeHTML.captured = []
    monkeypatch.setattr(weasyprint, 'HTML', _FakeHTML)

    resp = auth_client.get('/centros/export/pdf')
    assert resp.status_code == 200
    html = _FakeHTML.captured[0]
    assert '<script>' not in html, 'HTML sin escapar en PDF'
    assert '&lt;script&gt;' in html, 'escape ausente en PDF'
