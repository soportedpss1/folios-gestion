"""Importar folios desde .xlsx para marcar digitado (vista previa + confirmación)."""

import json
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models.audit_log import AuditLog
from app.models.folio import Folio
from tests.conftest import make_folio

HEADERS = ['folio']
Y = 2026


def make_xlsx(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def post_importar(client, rows, nombre='import.xlsx'):
    return client.post(
        '/folios/importar',
        data={'archivo': (make_xlsx(HEADERS, rows), nombre)},
        content_type='multipart/form-data',
    )


def confirmar(client, numeros):
    return client.post(
        '/folios/importar/confirmar',
        data={'payload': json.dumps(numeros)},
        follow_redirects=True,
    )


def test_importar_requiere_login(client):
    assert client.get('/folios/importar').status_code == 302
    assert client.post('/folios/importar/confirmar').status_code == 302
    assert client.get('/folios/plantilla').status_code == 302


def test_plantilla_es_xlsx_con_headers(admin_client):
    resp = admin_client.get('/folios/plantilla')
    assert resp.status_code == 200
    assert 'spreadsheetml' in resp.content_type
    wb = load_workbook(BytesIO(resp.data))
    assert [c.value for c in wb.active[1]] == HEADERS


def test_importar_bloqueado_para_operador(auth_client):
    resp = auth_client.get('/folios/importar', follow_redirects=True)
    assert resp.status_code == 200
    assert 'No tiene permisos' in resp.get_data(as_text=True)
    assert '/dashboard' in resp.request.path


def test_importar_bloqueado_para_lectura(client, app):
    from tests.conftest import make_user
    make_user('lectura', 'lectura')
    client.post('/auth/login', data={'username': 'lectura', 'password': 'pass12345'})
    resp = client.get('/folios/importar', follow_redirects=True)
    assert 'No tiene permisos' in resp.get_data(as_text=True)


def test_boton_importar_oculto_para_operador(auth_client):
    assert '/folios/importar' not in auth_client.get('/folios/').get_data(as_text=True)


def test_boton_importar_visible_para_admin(admin_client):
    assert '/folios/importar' in admin_client.get('/folios/').get_data(as_text=True)


def test_preview_coincidencia(admin_client, tipo, folio):
    resp = post_importar(admin_client, [[1000]])
    text = resp.get_data(as_text=True)
    assert '<strong>1</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'Confirmar importación' in text
    assert db.session.get(Folio, folio.id).digitado is False, \
        'la vista previa no debe marcar nada'


def test_preview_no_existente(admin_client):
    resp = post_importar(admin_client, [[9999]])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'No existe' in text
    assert 'Confirmar importación' not in text


def test_preview_numero_invalido(admin_client):
    resp = post_importar(admin_client, [['no-es-numero']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> folio(s) válido(s)' in text
    assert 'Número inválido' in text


def test_preview_repetido_entre_filas(admin_client, tipo, folio):
    resp = post_importar(admin_client, [[1000], [1000]])
    text = resp.get_data(as_text=True)
    assert '<strong>1</strong> folio(s) válido(s) de <strong>2</strong>' in text
    assert 'Repetido en el archivo' in text


def test_preview_numero_matchea_varios_folios(admin_client, tipo):
    """Mismo número entre año/tipo: el preview muestra todas las coincidencias."""
    make_folio(tipo, anio=Y, numero=1000, rangoId=1)
    make_folio(tipo, anio=Y - 1, numero=1000, rangoId=2)
    db.session.commit()
    resp = post_importar(admin_client, [[1000]])
    text = resp.get_data(as_text=True)
    assert '<strong>1</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert '<td>2</td>' in text, 'coincidencias = 2'


def test_preview_ya_digitado(admin_client, tipo):
    make_folio(tipo, anio=Y, numero=1000, digitado=True)
    resp = post_importar(admin_client, [[1000]])
    text = resp.get_data(as_text=True)
    assert '<strong>1</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'Ya digitado' in text


def test_confirm_marca_digitado(admin_client, tipo, folio):
    resp = confirmar(admin_client, [1000])

    assert resp.status_code == 200
    assert db.session.get(Folio, folio.id).digitado is True
    log = AuditLog.query.filter_by(tabla='folios', registro_id=folio.id).one()
    assert log.accion == 'UPDATE'
    text = resp.get_data(as_text=True)
    assert '1 folio marcado como digitado' in text


def test_confirm_marca_todas_las_coincidencias(admin_client, tipo):
    f1 = make_folio(tipo, anio=Y, numero=1000, rangoId=1)
    f2 = make_folio(tipo, anio=Y - 1, numero=1000, rangoId=2)
    db.session.commit()

    resp = confirmar(admin_client, [1000])

    assert '2 folios marcados como digitado' in resp.get_data(as_text=True)
    assert db.session.get(Folio, f1.id).digitado is True
    assert db.session.get(Folio, f2.id).digitado is True


def test_confirm_omite_ya_digitado(admin_client, tipo, folio):
    folio.digitado = True
    db.session.commit()

    resp = confirmar(admin_client, [1000])

    assert 'todos ya estaban digitados' in resp.get_data(as_text=True)
    assert AuditLog.query.filter_by(tabla='folios').count() == 0, \
        'sin cambios no debe auditar'


def test_confirm_numero_inexistente_omite(admin_client):
    resp = confirmar(admin_client, [9999])
    assert resp.status_code == 200
    assert 'todos ya estaban digitados' in resp.get_data(as_text=True)


def test_confirm_payload_invalido_no_marca_nada(admin_client, folio):
    resp = admin_client.post('/folios/importar/confirmar',
                             data={'payload': 'no-es-json'},
                             follow_redirects=True)
    assert 'inválidos' in resp.get_data(as_text=True)
    assert db.session.get(Folio, folio.id).digitado is False


def test_confirm_payload_con_bool_no_marca_nada(admin_client, folio):
    resp = admin_client.post('/folios/importar/confirmar',
                             data={'payload': json.dumps([True])},
                             follow_redirects=True)
    assert 'inválidos' in resp.get_data(as_text=True)
    assert db.session.get(Folio, folio.id).digitado is False


def test_archivo_csv_rechazado(admin_client):
    resp = admin_client.post(
        '/folios/importar',
        data={'archivo': (BytesIO(b'folio\n1000\n'), 'import.csv')},
        content_type='multipart/form-data',
    )
    assert resp.status_code == 200
    assert 'debe ser un Excel' in resp.get_data(as_text=True)


def test_headers_incorrectos_rechazados(admin_client):
    resp = admin_client.post(
        '/folios/importar',
        data={'archivo': (make_xlsx(['otra'], [[1000]]), 'import.xlsx')},
        content_type='multipart/form-data',
    )
    text = resp.get_data(as_text=True)
    assert 'Encabezados incorrectos' in text
