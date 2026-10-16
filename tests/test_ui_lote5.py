"""Lote 5 de la auditoría UI/UX: consistencia — estados vacíos unificados,
columnas compartidas entre PDF y Excel de reportes, botones de export
visibles, textos y toasts homogéneos, agrupación de botones del índice de
folios y guards JS contra carreras de fetch + CSS muerto fuera."""

import re
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.extensions import db
from tests.conftest import make_folio

RAIZ = Path(__file__).resolve().parents[1]
PLANTILLAS = RAIZ / 'app' / 'templates'

COLUMNAS_REPORTE = [
    'Folio', 'Tipo Certificado', 'Año', 'Centro', 'Estado',
    'Digitado', 'Escaneado', 'Nulo',
    'Fecha Entrega', 'Entregado por',
    'Fecha Devolución', 'Recibido Por',
]


def _tag_attr(body, patrón):
    match = re.search(patrón, body, re.S)
    return match.group(0) if match else None


def _capturar_pdf(monkeypatch):
    try:
        import weasyprint
    except OSError as exc:
        pytest.skip(f'weasyprint sin libpango (solo Docker): {exc}')

    class _FakeHTML:
        captured = []

        def __init__(self, string=None, **kwargs):
            _FakeHTML.captured.append(string)

        def write_pdf(self):
            return b'%PDF-1.4 fake'

    _FakeHTML.captured = []
    monkeypatch.setattr(weasyprint, 'HTML', _FakeHTML)
    return _FakeHTML


# --- 1. Estados vacíos: sin {% else %} en usuarios/permisos; textos unidos ---

def test_estados_vacios_usuarios_y_permisos_en_plantillas():
    usuarios = (PLANTILLAS / 'usuarios' / 'index.html').read_text(encoding='utf-8')
    assert 'No hay usuarios registrados' in usuarios, \
        'índice de usuarios sin fila de estado vacío (tbody en blanco)'
    assert 'colspan="7"' in usuarios, 'la fila vacía no cubre las 7 columnas'

    permisos = (PLANTILLAS / 'permisos' / 'index.html').read_text(encoding='utf-8')
    assert 'No hay usuarios registrados' in permisos, \
        'índice de permisos sin fila de estado vacío'
    assert 'colspan="6"' in permisos, 'la fila vacía no cubre las 6 columnas'


@pytest.mark.parametrize('ruta, texto, viejo', [
    ('/folios/', 'No hay folios registrados', 'No se encontraron folios'),
    ('/folios/transferencias', 'No hay transferencias registradas',
     'No se encontraron transferencias'),
])
def test_estados_vacios_texto_unificado(admin_client, ruta, texto, viejo):
    body = admin_client.get(ruta).get_data(as_text=True)
    assert texto in body, f'{ruta}: estado vacío fuera del patrón "No hay …"'
    assert viejo not in body, f'{ruta}: patrón antiguo "{viejo}"'


# --- 2. Reportes: PDF y Excel con las mismas columnas en el mismo orden ---

def test_reportes_excel_con_columnas_canonicas(admin_client, app, tipo):
    make_folio(tipo, numero=6050, digitado=True)
    db.session.commit()

    resp = admin_client.get('/reportes/export/excel')
    assert resp.status_code == 200
    filas = list(load_workbook(BytesIO(resp.data)).active.iter_rows(values_only=True))
    assert list(filas[0]) == COLUMNAS_REPORTE, \
        'el Excel de reportes no usa el orden de columnas compartido'
    assert filas[1][5] == 'Sí', 'digitado sin tilde'


def test_reportes_pdf_comparte_columnas_con_excel(admin_client, app, tipo,
                                                  monkeypatch):
    fake = _capturar_pdf(monkeypatch)
    make_folio(tipo, numero=6050, digitado=True)
    db.session.commit()

    resp = admin_client.get('/reportes/export/pdf')
    assert resp.status_code == 200
    html = fake.captured[0]

    theads = re.findall(r'<th>([^<]+)</th>', html)
    assert theads == COLUMNAS_REPORTE, \
        f'columnas del PDF {theads} ≠ columnas compartidas'
    assert html.count('<td>') == 12, 'una fila debe tener 12 celdas (12 columnas)'
    assert 'Si</td>' not in html, '"Si" sin tilde en el PDF'
    assert re.search(r'<td>Sí</td>', html), '"Sí" sin tilde en el PDF'

    xl = admin_client.get('/reportes/export/excel')
    encabezado_excel = list(
        load_workbook(BytesIO(xl.data)).active.iter_rows(values_only=True))[0]
    assert list(theads) == list(encabezado_excel), \
        'el PDF y el Excel del mismo reporte no comparten columnas'


# --- 3. Botones de export de reportes visibles y etiquetados como el resto ---

def test_botones_export_visibles_y_nombrados(admin_client):
    body = admin_client.get('/reportes/').get_data(as_text=True)
    for ident, esperado in (('exportExcel', '/reportes/export/excel'),
                            ('exportPdf', '/reportes/export/pdf')):
        tag = _tag_attr(body, rf'<a[^>]*id="{ident}"[^>]*>')
        assert tag, f'falta #{ident}'
        assert 'display:none' not in tag, \
            f'#{ident} oculto hasta que JS decida mostrarlo'
        assert f'href="{esperado}' in tag, \
            f'#{ident} sin href del server (queda "#" si JS no corre)'
    assert 'Exportar Excel' not in body and 'Exportar PDF' not in body, \
        'etiquetas distintas a "Excel"/"PDF" del resto de módulos'


# --- 4. Textos homogéneos ---

def test_reportes_usa_clic(admin_client):
    body = admin_client.get('/reportes/').get_data(as_text=True)
    assert 'haga clic' in body and 'haga click' not in body


def test_permisos_indice_sin_grants_ingles(admin_client):
    body = admin_client.get('/permisos/').get_data(as_text=True)
    assert '<th>Permisos</th>' in body, 'la columna sigue llamándose "Grants"'
    assert 'Grants' not in body


def test_detalles_folio_sin_jerga(admin_client, folio):
    body = admin_client.get(f'/folios/{folio.id}').get_data(as_text=True)
    assert 'Rango (id interno)' in body, '"Rango ID" crudo para el usuario'
    assert 'Rango ID' not in body
    assert '>N/A<' not in body, 'los huecos deben ser "-" como en el resto'


def test_flashes_sin_correctamente():
    culpables = []
    for archivo in Path(RAIZ, 'app').rglob('*.py'):
        for num, linea in enumerate(
                archivo.read_text(encoding='utf-8').splitlines(), 1):
            if 'flash(' in linea and 'correctamente' in linea:
                culpables.append(f'{archivo.relative_to(RAIZ)}:{num}')
    assert not culpables, f'toasts sin homogeneizar a "exitosamente": {culpables}'


# --- 5. Índice de folios: botones agrupados (no esparcidos por space-between) ---

def test_folios_header_botones_agrupados():
    src = (PLANTILLAS / 'folios' / 'index.html').read_text(encoding='utf-8')
    assert '<div class="d-flex gap-2 flex-wrap">' in src, \
        'los botones del encabezado no están en un contenedor con gap'
    assert 'btn-outline-danger ms-2' not in src and \
        'btn-outline-info ms-2' not in src, 'ms-2 suelto en vez de gap'


# --- 6. Guards JS: carreras de fetch y tooltips ---

@pytest.mark.parametrize('rel', ['entregas/create.html', 'devoluciones/create.html'])
def test_loadfolios_guard_de_carrera(rel):
    src = (PLANTILLAS / rel).read_text(encoding='utf-8')
    assert 't !== seq' in src, \
        f'{rel}: respuestas fuera de orden pueden pintar datos viejos'


def test_loaddata_reportes_guard_de_carrera():
    src = (RAIZ / 'app' / 'static' / 'js' / 'reports.js').read_text(encoding='utf-8')
    assert 't !== seq' in src, \
        'loadData sin guard: clics rápidos dejan la página vieja en pantalla'


def test_entregas_error_http_visible():
    src = (PLANTILLAS / 'entregas' / 'create.html').read_text(encoding='utf-8')
    assert 'Error al cargar los folios' in src, \
        'loadFolios devuelve silenciosamente si !resp.ok'


def test_main_js_tooltip_sync_timer_sin_congelar():
    src = (RAIZ / 'app' / 'static' / 'js' / 'main.js').read_text(encoding='utf-8')
    assert "el.id === 'sync-timer'" in src, \
        'el tooltip de JS se queda con el title viejo aunque pintar() lo actualice'


# --- 7. CSS muerto fuera ---

def test_css_muerto_fuera():
    main = (RAIZ / 'app' / 'static' / 'css' / 'main.css').read_text(encoding='utf-8')
    variables = (RAIZ / 'app' / 'static' / 'css' / 'variables.css').read_text(
        encoding='utf-8')

    assert 'skeleton' not in main.lower(), '.skeleton sin uso en ninguna plantilla'
    assert 'skeleton' not in variables.lower()
    assert 'text-truncate-2' not in main, '.text-truncate-2 sin uso'
    for var in ('--primary-darker', '--accent-rgb', '--sidebar-collapsed-width'):
        assert var not in variables, f'{var} declarado y nunca usado'
    assert main.count('@media (prefers-reduced-motion: reduce)') == 1, \
        'bloque reduced-motion redundante (el scroll no tiene transiciones)'
    assert '.modal-header :focus-visible' not in main, \
        'anillo blanco sobre modal-header claro: el foco queda invisible'
