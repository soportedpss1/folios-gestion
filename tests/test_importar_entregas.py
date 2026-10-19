"""Importar entregas de folios desde .xlsx (vista previa + confirmación).

La fecha va por fila en el Excel (columna obligatoria `fecha`): no hay campo
de fecha en el formulario ni hidden en la confirmación. Cada movimiento
guarda la fecha de su fila.

El centro de salud TAMBIÉN viene por fila: la columna `centro` lleva el
nombre exacto del centro en la DB. No hay select de centro en el form ni
hidden en la confirmación; un archivo puede mezclar centros distintos.

TODAS las columnas son obligatorias (incluida la final `usuario`, con el
username del responsable): un archivo sin ella se rechaza y una celda
vacía deja la fila inválida en la vista previa.
"""

import json
from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models.centro import Centro
from app.models.entrega_folio import EntregaFolio
from app.models.folio import Folio
from tests.conftest import make_folio

HEADERS = ['folio', 'tipoCert', 'anioCert', 'fecha', 'centro', 'usuario']
Y = 2026


def make_xlsx(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def post_importar(client, rows, nombre='import.xlsx'):
    return client.post(
        '/entregas/importar',
        data={'archivo': (make_xlsx(HEADERS, rows), nombre)},
        content_type='multipart/form-data',
    )


def confirmar(client, items):
    """items: lista de {'id', 'fecha', 'centroId', ['usuarioId']}."""
    return client.post(
        '/entregas/importar/confirmar',
        data={'payload': json.dumps(items)},
        follow_redirects=True,
    )


def test_importar_requiere_login(client):
    assert client.get('/entregas/importar').status_code == 302
    assert client.post('/entregas/importar/confirmar').status_code == 302


def test_plantilla_es_xlsx_con_headers(admin_client, centro):
    resp = admin_client.get('/entregas/plantilla')
    assert resp.status_code == 200
    assert 'spreadsheetml' in resp.content_type
    wb = load_workbook(BytesIO(resp.data))
    assert [c.value for c in wb.active[1]] == HEADERS, \
        'la plantilla lleva todas las columnas obligatorias (usuario al final)'
    ejemplo = [c.value for c in wb.active[2]]
    assert ejemplo[4] == centro.name, 'la fila de ejemplo lleva un centro real'


def test_formulario_no_tiene_campo_fecha(admin_client, centro):
    text = admin_client.get('/entregas/importar').get_data(as_text=True)
    assert 'id="fechaEntrega"' not in text, 'la fecha se toma del Excel'


def test_formulario_no_tiene_selector_de_centro(admin_client):
    text = admin_client.get('/entregas/importar').get_data(as_text=True)
    assert 'id="centroId"' not in text, 'el centro se toma del Excel'


def test_preview_resuelve_disponible(admin_client, tipo, centro):
    make_folio(tipo, anio=Y, numero=1000)
    resp = post_importar(admin_client,
                         [[1000, tipo.name, Y, '01/10/2026', centro.name,
                           'adminuser']])
    text = resp.get_data(as_text=True)
    assert '<strong>1</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'Confirmar importación' in text
    assert 'Todas las filas del archivo son válidas' in text, \
        'sin inválidas no se lista ninguna fila (la fecha por fila se verifica en la lista de inválidas)'
    assert '<input type="hidden" name="fecha"' not in text, 'sin hidden de fecha'
    assert 'id="centroId"' not in text, 'sin hidden/select de centro'
    assert EntregaFolio.query.count() == 0, 'la vista previa no debe crear nada'


def test_preview_marca_no_disponible(admin_client, tipo, centro):
    make_folio(tipo, anio=Y, numero=1000, estado='entregado')
    resp = post_importar(admin_client,
                         [[1000, tipo.name, Y, '2026-10-01', centro.name,
                           'adminuser']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'no está disponible' in text


def test_preview_marca_inexistente(admin_client, tipo, centro):
    resp = post_importar(admin_client,
                         [[9999, tipo.name, Y, '2026-10-01', centro.name,
                           'adminuser']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'no existe' in text


def test_preview_lista_solo_filas_invalidas(admin_client, tipo, centro):
    make_folio(tipo, anio=Y, numero=1000)
    resp = post_importar(admin_client, [
        [1000, tipo.name, Y, '01/10/2026', centro.name, 'adminuser'],
        [9999, tipo.name, Y, '01/10/2026', centro.name, 'adminuser'],
    ])
    text = resp.get_data(as_text=True)
    assert '<strong>1</strong> folio(s) válido(s) de <strong>2</strong>' in text, \
        'el resumen sigue contando válidas y total'
    assert 'El folio no existe.' in text, 'la fila inválida se lista con su error'
    assert '01/10/2026' in text, 'la fila inválida muestra su fecha por fila'
    assert 'Folio 1000' not in text, \
        'la lista de validación no debe mostrar filas válidas'


def test_preview_sin_filas_invalidas_muestra_mensaje(admin_client, tipo, centro):
    make_folio(tipo, anio=Y, numero=1000)
    resp = post_importar(admin_client,
                         [[1000, tipo.name, Y, '01/10/2026', centro.name,
                           'adminuser']])
    text = resp.get_data(as_text=True)
    assert 'Todas las filas del archivo son válidas' in text, \
        'sin inválidas se muestra el mensaje en lugar de la tabla'
    assert 'Folio 1000' not in text, 'no debe haber filas en la lista'
    assert 'Confirmar importación' in text, 'el resumen y la confirmación siguen'


def test_preview_rechaza_archivo_sin_columna_fecha(admin_client, tipo, centro):
    resp = admin_client.post('/entregas/importar', data={
        'archivo': (make_xlsx(['folio', 'tipoCert', 'anioCert', 'centro'],
                              [[1000, tipo.name, Y, centro.name]]),
                    'sin-fecha.xlsx'),
    }, content_type='multipart/form-data')
    assert 'Encabezados incorrectos' in resp.get_data(as_text=True)


def test_preview_rechaza_archivo_sin_columna_centro(admin_client, tipo):
    resp = admin_client.post('/entregas/importar', data={
        'archivo': (make_xlsx(['folio', 'tipoCert', 'anioCert', 'fecha'],
                              [[1000, tipo.name, Y, '2026-10-01']]),
                    'sin-centro.xlsx'),
    }, content_type='multipart/form-data')
    assert 'Encabezados incorrectos' in resp.get_data(as_text=True)


def test_preview_rechaza_archivo_sin_columna_usuario(admin_client, tipo, centro):
    """Todas las columnas son obligatorias: sin `usuario` no hay importación."""
    resp = admin_client.post('/entregas/importar', data={
        'archivo': (make_xlsx(HEADERS[:-1],
                              [[1000, tipo.name, Y, '2026-10-01', centro.name]]),
                    'sin-usuario.xlsx'),
    }, content_type='multipart/form-data')
    assert 'Encabezados incorrectos' in resp.get_data(as_text=True)


def test_preview_celda_usuario_vacia_marca_fila(admin_client, tipo, centro):
    """La celda también es obligatoria: vacía → fila inválida."""
    make_folio(tipo, anio=Y, numero=1000)
    resp = post_importar(admin_client,
                         [[1000, tipo.name, Y, '2026-10-01', centro.name, '']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'Usuario vacío.' in text


def test_preview_marca_fecha_invalida(admin_client, tipo, centro):
    make_folio(tipo, anio=Y, numero=1000)
    resp = post_importar(admin_client,
                         [[1000, tipo.name, Y, 'ayer', centro.name,
                           'adminuser']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'Fecha inválida' in text


def test_preview_centro_desconocido_marca_fila(admin_client, tipo, centro):
    make_folio(tipo, anio=Y, numero=1000)
    resp = post_importar(admin_client,
                         [[1000, tipo.name, Y, '2026-10-01', 'No Existe',
                           'adminuser']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'Centro de salud desconocido: No Existe.' in text


def test_preview_centro_es_sensible_a_mayusculas(admin_client, tipo, centro):
    """El match del nombre es EXACTO: 'centro de prueba' no es 'Centro de Prueba'."""
    make_folio(tipo, anio=Y, numero=1000)
    resp = post_importar(admin_client,
                         [[1000, tipo.name, Y, '2026-10-01',
                           centro.name.lower(), 'adminuser']])
    assert 'Centro de salud desconocido' in resp.get_data(as_text=True)


def test_preview_centro_vacio_marca_fila(admin_client, tipo, centro):
    make_folio(tipo, anio=Y, numero=1000)
    resp = post_importar(admin_client,
                         [[1000, tipo.name, Y, '2026-10-01', '',
                           'adminuser']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'Centro de salud vacío.' in text


def test_preview_centro_inactivo_marca_fila(admin_client, tipo, centro):
    baja = Centro(name='Centro Baja', telefono='555', direccion='Calle 2',
                  activo=False)
    db.session.add(baja)
    db.session.commit()
    make_folio(tipo, anio=Y, numero=1000)

    resp = post_importar(admin_client,
                         [[1000, tipo.name, Y, '2026-10-01', 'Centro Baja',
                           'adminuser']])
    assert 'Centro de salud inactivo: Centro Baja.' in \
        resp.get_data(as_text=True)


def test_preview_admite_centros_mezclados(admin_client, tipo, centro):
    otro = Centro(name='Otro Centro', telefono='555', direccion='Calle 1',
                  activo=True)
    db.session.add(otro)
    db.session.commit()
    make_folio(tipo, anio=Y, numero=1000)
    make_folio(tipo, anio=Y, numero=1001)

    resp = post_importar(admin_client, [
        [1000, tipo.name, Y, '2026-10-01', centro.name, 'adminuser'],
        [1001, tipo.name, Y, '2026-10-01', otro.name, 'adminuser'],
    ])
    text = resp.get_data(as_text=True)
    assert '<strong>2</strong> folio(s) válido(s) de <strong>2</strong>' in text
    assert f'&#34;centroId&#34;: {centro.id}' in text, 'payload con centro por fila'
    assert f'&#34;centroId&#34;: {otro.id}' in text


def test_confirm_registra_entrega(admin_client, tipo, centro):
    folio = make_folio(tipo, anio=Y, numero=1000)
    resp = confirmar(admin_client, [
        {'id': folio.id, 'fecha': '2026-10-01', 'centroId': centro.id},
    ])

    assert resp.status_code == 200
    assert EntregaFolio.query.count() == 1
    assert db.session.get(Folio, folio.id).estado == 'entregado'
    entrega = EntregaFolio.query.one()
    assert entrega.fechaEntrega == date(2026, 10, 1)
    assert entrega.centroId == centro.id
    assert '1 entrega registrada' in resp.get_data(as_text=True)


def test_confirm_respeta_centro_distinto_por_fila(admin_client, tipo, centro):
    otro = Centro(name='Otro Centro', telefono='555', direccion='Calle 1',
                  activo=True)
    db.session.add(otro)
    db.session.commit()
    f1 = make_folio(tipo, anio=Y, numero=1000)
    f2 = make_folio(tipo, anio=Y, numero=1001)

    resp = confirmar(admin_client, [
        {'id': f1.id, 'fecha': '2026-10-01', 'centroId': centro.id},
        {'id': f2.id, 'fecha': '2026-10-01', 'centroId': otro.id},
    ])

    assert resp.status_code == 200
    assert EntregaFolio.query.count() == 2
    assert {e.centroId for e in EntregaFolio.query.all()} == \
        {centro.id, otro.id}, 'cada folio queda en el centro de su fila'


def test_confirm_respeta_fecha_distinta_por_fila(admin_client, tipo, centro):
    f1 = make_folio(tipo, anio=Y, numero=1000)
    f2 = make_folio(tipo, anio=Y, numero=1001)
    resp = confirmar(admin_client, [
        {'id': f1.id, 'fecha': '2026-10-01', 'centroId': centro.id},
        {'id': f2.id, 'fecha': '05/10/2026', 'centroId': centro.id},
    ])

    assert resp.status_code == 200
    assert EntregaFolio.query.count() == 2
    assert {e.fechaEntrega for e in EntregaFolio.query.all()} == \
        {date(2026, 10, 1), date(2026, 10, 5)}


def test_confirm_revalida_estado(admin_client, tipo, centro):
    """Folio que dejó de estar disponible entre previa y confirm → omitido."""
    folio = make_folio(tipo, anio=Y, numero=1000, estado='entregado')
    resp = confirmar(admin_client, [
        {'id': folio.id, 'fecha': '2026-10-01', 'centroId': centro.id},
    ])

    assert EntregaFolio.query.count() == 0
    assert 'Ningún folio sigue disponible' in resp.get_data(as_text=True)


def test_confirm_payload_invalido_no_crea_nada(admin_client, tipo, centro):
    resp = admin_client.post('/entregas/importar/confirmar', data={
        'payload': 'no-es-json',
    }, follow_redirects=True)

    assert 'inválidos' in resp.get_data(as_text=True)
    assert EntregaFolio.query.count() == 0


def test_confirm_rechaza_fecha_invalida_en_payload(admin_client, tipo, centro):
    """El hidden del confirm es editable: fecha faltante o basura → rechazo."""
    folio = make_folio(tipo, anio=Y, numero=1000)
    resp = confirmar(admin_client, [
        {'id': folio.id, 'fecha': 'no-fecha', 'centroId': centro.id},
    ])

    assert 'inválidos' in resp.get_data(as_text=True)
    assert EntregaFolio.query.count() == 0


def test_confirm_rechaza_centro_ausente_en_payload(admin_client, tipo, centro):
    """Sin `centroId` por fila el payload completo es inválido."""
    folio = make_folio(tipo, anio=Y, numero=1000)
    resp = confirmar(admin_client, [
        {'id': folio.id, 'fecha': '2026-10-01'},
    ])

    assert 'inválidos' in resp.get_data(as_text=True)
    assert EntregaFolio.query.count() == 0


def test_confirm_rechaza_centro_desconocido(admin_client, tipo, centro):
    """El hidden es editable: centro inexistente o inactivo → nada se registra."""
    folio = make_folio(tipo, anio=Y, numero=1000)
    resp = confirmar(admin_client, [
        {'id': folio.id, 'fecha': '2026-10-01', 'centroId': 999999},
    ])

    assert 'Centro de salud inválido.' in resp.get_data(as_text=True)
    assert EntregaFolio.query.count() == 0
    assert db.session.get(Folio, folio.id).estado == 'disponible'


def test_confirm_rechaza_centro_inactivo(admin_client, tipo, centro):
    baja = Centro(name='Centro Baja', telefono='555', direccion='Calle 2',
                  activo=False)
    db.session.add(baja)
    db.session.commit()
    folio = make_folio(tipo, anio=Y, numero=1000)

    resp = confirmar(admin_client, [
        {'id': folio.id, 'fecha': '2026-10-01', 'centroId': baja.id},
    ])

    assert 'Centro de salud inválido.' in resp.get_data(as_text=True)
    assert EntregaFolio.query.count() == 0
