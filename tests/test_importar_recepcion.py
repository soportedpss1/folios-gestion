"""Importar rangos de recepción desde .xlsx (vista previa + confirmación)."""

import json
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.models.folio import Folio
from app.models.recepcion_folio import RecepcionFolio

HEADERS = ['fecha', 'anioCert', 'tipoCert', 'folioInicial', 'folioFinal']


def make_xlsx(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def post_importar(client, headers, rows, nombre='import.xlsx'):
    return client.post(
        '/recepcion/importar',
        data={'archivo': (make_xlsx(headers, rows), nombre)},
        content_type='multipart/form-data',
    )


def test_importar_requiere_login(client):
    assert client.get('/recepcion/importar').status_code == 302
    assert client.post('/recepcion/importar/confirmar').status_code == 302


def test_plantilla_requiere_login(client):
    assert client.get('/recepcion/plantilla').status_code == 302


def test_plantilla_es_xlsx_con_headers(admin_client):
    resp = admin_client.get('/recepcion/plantilla')
    assert resp.status_code == 200
    assert 'spreadsheetml' in resp.content_type
    wb = load_workbook(BytesIO(resp.data))
    assert [c.value for c in wb.active[1]] == HEADERS + ['usuario'], \
        'la plantilla trae la columna usuario final (opcional al importar)'


def test_importar_sin_permisos_redirige(client, app):
    from tests.conftest import make_user
    make_user('lectura', 'lectura')
    client.post('/auth/login', data={'username': 'lectura', 'password': 'pass12345'})
    resp = client.get('/recepcion/importar', follow_redirects=True)
    assert resp.status_code == 200
    assert 'No tiene permisos' in resp.get_data(as_text=True)


def test_preview_muestra_fila_valida(admin_client, tipo):
    resp = post_importar(admin_client, HEADERS, [['2026-10-01', 2026, tipo.name, 1, 5]])
    assert resp.status_code == 200
    text = resp.get_data(as_text=True)
    assert '<strong>1</strong> fila(s) válida(s) de <strong>1</strong>' in text
    assert 'Confirmar importación' in text
    assert RecepcionFolio.query.count() == 0, 'la vista previa no debe crear nada'


def test_preview_marca_solape_contra_db(admin_client, tipo):
    admin_client.post('/recepcion/create', data={
        'fecha': '2026-09-28', 'anioCert': 2026, 'tipoCert': tipo.id,
        'folioInicial': 100, 'folioFinal': 110,
    }, follow_redirects=True)

    resp = post_importar(admin_client, HEADERS, [['2026-10-01', 2026, tipo.name, 105, 120]])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> fila(s) válida(s) de <strong>1</strong>' in text
    assert 'solapa' in text


def test_preview_marca_solape_entre_filas_del_archivo(admin_client, tipo):
    resp = post_importar(admin_client, HEADERS, [
        ['2026-10-01', 2026, tipo.name, 1, 10],
        ['2026-10-02', 2026, tipo.name, 5, 15],
    ])
    text = resp.get_data(as_text=True)
    assert '<strong>1</strong> fila(s) válida(s) de <strong>2</strong>' in text
    assert 'otra fila del archivo' in text


def test_confirm_importa_rangos(admin_client, tipo):
    payload = json.dumps([{
        'fecha': '2026-10-01', 'anioCert': 2026, 'tipoCert_id': tipo.id,
        'folioInicial': 1, 'folioFinal': 5,
    }])
    resp = admin_client.post('/recepcion/importar/confirmar',
                            data={'payload': payload}, follow_redirects=True)

    assert resp.status_code == 200
    assert RecepcionFolio.query.count() == 1
    assert Folio.query.count() == 5
    text = resp.get_data(as_text=True).lower()
    assert '1 recepción registrada, 5 folios creados' in text


def test_confirm_payload_invalido_no_crea_nada(admin_client, tipo):
    resp = admin_client.post('/recepcion/importar/confirmar',
                            data={'payload': 'no-es-json'}, follow_redirects=True)
    assert 'inválidos' in resp.get_data(as_text=True)
    assert RecepcionFolio.query.count() == 0


def test_confirm_rango_sobre_maximo_rechazado(admin_client, tipo):
    payload = json.dumps([{
        'fecha': '2026-10-01', 'anioCert': 2026, 'tipoCert_id': tipo.id,
        'folioInicial': 1, 'folioFinal': 99999,
    }])
    admin_client.post('/recepcion/importar/confirmar',
                     data={'payload': payload}, follow_redirects=True)
    assert RecepcionFolio.query.count() == 0


def test_archivo_csv_rechazado(admin_client):
    resp = admin_client.post(
        '/recepcion/importar',
        data={'archivo': (BytesIO(b'a,b,c\n1,2,3\n'), 'import.csv')},
        content_type='multipart/form-data',
    )
    assert resp.status_code == 200
    assert 'debe ser un Excel' in resp.get_data(as_text=True)


def test_headers_incorrectos_rechazados(admin_client):
    resp = post_importar(admin_client, ['otra', 'columna'], [[1, 2]])
    text = resp.get_data(as_text=True)
    assert 'Encabezados incorrectos' in text
    assert RecepcionFolio.query.count() == 0
