"""Importar centros desde .xlsx (vista previa + confirmación)."""

import json
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.models.centro import Centro
from app.models.audit_log import AuditLog

HEADERS = ['nombre', 'telefono', 'direccion']


def make_xlsx(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def post_importar(client, headers, rows, nombre='import.xlsx'):
    return client.post(
        '/centros/importar',
        data={'archivo': (make_xlsx(headers, rows), nombre)},
        content_type='multipart/form-data',
    )


def test_importar_requiere_login(client):
    assert client.get('/centros/importar').status_code == 302
    assert client.post('/centros/importar/confirmar').status_code == 302


def test_plantilla_requiere_login(client):
    assert client.get('/centros/plantilla').status_code == 302


def test_plantilla_es_xlsx_con_headers(admin_client):
    resp = admin_client.get('/centros/plantilla')
    assert resp.status_code == 200
    assert 'spreadsheetml' in resp.content_type
    wb = load_workbook(BytesIO(resp.data))
    assert [c.value for c in wb.active[1]] == HEADERS


def test_importar_sin_permisos_redirige(client, app):
    from tests.conftest import make_user
    make_user('lectura', 'lectura')
    client.post('/auth/login', data={'username': 'lectura', 'password': 'pass12345'})
    resp = client.get('/centros/importar', follow_redirects=True)
    assert resp.status_code == 200
    assert 'No tiene permisos' in resp.get_data(as_text=True)


def test_preview_muestra_fila_valida(admin_client):
    resp = post_importar(admin_client, HEADERS,
                         [['Hospital Norte', '555555', 'Calle 1']])
    assert resp.status_code == 200
    text = resp.get_data(as_text=True)
    assert '<strong>1</strong> fila(s) válida(s) de <strong>1</strong>' in text
    assert 'Confirmar importación' in text
    assert Centro.query.count() == 0, 'la vista previa no debe crear nada'


def test_preview_marca_duplicado_contra_db(admin_client, centro):
    resp = post_importar(admin_client, HEADERS,
                         [[centro.name, '555555', 'Otra dirección']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> fila(s) válida(s) de <strong>1</strong>' in text
    assert 'ya existe' in text


def test_preview_marca_duplicado_entre_filas(admin_client):
    resp = post_importar(admin_client, HEADERS, [
        ['Hospital Norte', '555555', 'Calle 1'],
        ['hospital norte', '555556', 'Calle 2'],
    ])
    text = resp.get_data(as_text=True)
    assert '<strong>1</strong> fila(s) válida(s) de <strong>2</strong>' in text
    assert 'repetido en el archivo' in text


def test_preview_nombre_vacio_rechazado(admin_client):
    resp = post_importar(admin_client, HEADERS, [['  ', '555', 'Calle']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> fila(s) válida(s)' in text
    assert 'Nombre vacío' in text


def test_confirm_importa_centros(admin_client):
    payload = json.dumps([
        {'nombre': 'Hospital Norte', 'telefono': '555555', 'direccion': 'Calle 1'},
    ])
    resp = admin_client.post('/centros/importar/confirmar',
                             data={'payload': payload}, follow_redirects=True)

    assert resp.status_code == 200
    centro = Centro.query.filter_by(name='Hospital Norte').one()
    assert centro.activo is True
    assert centro.telefono == '555555'
    log = AuditLog.query.filter_by(tabla='centros', registro_id=centro.id).one()
    assert log.accion == 'INSERT'
    text = resp.get_data(as_text=True)
    assert '1 centro creado' in text


def test_confirm_rechaza_duplicado_por_payload_editado(admin_client, centro):
    payload = json.dumps([
        {'nombre': centro.name, 'telefono': '', 'direccion': ''},
    ])
    admin_client.post('/centros/importar/confirmar',
                      data={'payload': payload}, follow_redirects=True)
    assert Centro.query.count() == 1, 'el duplicado no debe insertarse'


def test_confirm_payload_invalido_no_crea_nada(admin_client):
    resp = admin_client.post('/centros/importar/confirmar',
                             data={'payload': 'no-es-json'}, follow_redirects=True)
    assert 'inválidos' in resp.get_data(as_text=True)
    assert Centro.query.count() == 0


def test_confirm_nombre_corto_rechazado(admin_client):
    payload = json.dumps([{'nombre': 'ab', 'telefono': '', 'direccion': ''}])
    admin_client.post('/centros/importar/confirmar',
                      data={'payload': payload}, follow_redirects=True)
    assert Centro.query.count() == 0


def test_archivo_csv_rechazado(admin_client):
    resp = admin_client.post(
        '/centros/importar',
        data={'archivo': (BytesIO(b'a,b,c\n1,2,3\n'), 'import.csv')},
        content_type='multipart/form-data',
    )
    assert resp.status_code == 200
    assert 'debe ser un Excel' in resp.get_data(as_text=True)


def test_headers_incorrectos_rechazados(admin_client):
    resp = post_importar(admin_client, ['otra', 'columna'], [[1, 2]])
    text = resp.get_data(as_text=True)
    assert 'Encabezados incorrectos' in text
    assert Centro.query.count() == 0
