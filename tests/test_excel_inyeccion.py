"""Fase 5 — inyección de fórmulas en Excel: openpyxl escribe '=...' como fórmula."""

from io import BytesIO

from openpyxl import load_workbook

from app.extensions import db
from app.models.entrega_folio import EntregaFolio
from app.models.recepcion_folio import RecepcionFolio
from datetime import date
from tests.conftest import make_folio

Y = 2026


def _celdas_formula(wb):
    return [
        f'{ws.title}!{c.coordinate}'
        for ws in wb.worksheets
        for fila in ws.iter_rows()
        for c in fila
        if c.data_type == 'f'
    ]


def test_reportes_excel_no_escribe_formulas(auth_client, app, tipo, centro):
    tipo.name = '=1+1'
    db.session.commit()
    make_folio(tipo, anio=Y, numero=7001)
    db.session.commit()

    resp = auth_client.get(f'/reportes/export/excel?anioCert={Y}')
    assert resp.status_code == 200

    wb = load_workbook(BytesIO(resp.data))
    formulas = _celdas_formula(wb)
    assert not formulas, f'celdas escritas como fórmula: {formulas}'
    # el contenido se conserva tal cual, solo cambia el tipo de celda
    assert wb.active.cell(row=2, column=2).value == '=1+1'


def test_centros_excel_no_escribe_formulas(auth_client, app, tipo, centro, operador):
    operador.username = '=HYPERLINK("http://evil","click")'
    db.session.commit()
    folio = make_folio(tipo, anio=Y, numero=7002, estado='entregado')
    db.session.add(EntregaFolio(folio_id=folio.id, fechaEntrega=date(Y, 6, 1),
                                centroId=centro.id, userId=operador.id))
    db.session.commit()

    resp = auth_client.get(f'/centros/{centro.id}/export/excel')
    assert resp.status_code == 200

    wb = load_workbook(BytesIO(resp.data))
    formulas = _celdas_formula(wb)
    assert not formulas, f'celdas escritas como fórmula: {formulas}'


def test_recepcion_excel_no_escribe_formulas(admin_client, app, tipo, operador):
    tipo.name = '=1+1'
    operador.username = '=HYPERLINK("http://evil","click")'
    db.session.add(RecepcionFolio(fecha=date(Y, 4, 1), anioCert=Y,
                                  tipoCert_id=tipo.id, folioInicial=400,
                                  folioFinal=410, rangoId=1, userId=operador.id))
    db.session.commit()

    resp = admin_client.get('/recepcion/export/excel')
    assert resp.status_code == 200

    wb = load_workbook(BytesIO(resp.data))
    formulas = _celdas_formula(wb)
    assert not formulas, f'celdas escritas como fórmula: {formulas}'
    assert wb.active.cell(row=2, column=4).value == '=1+1'
    assert wb.active.cell(row=2, column=8).value == operador.username
