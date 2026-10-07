"""Importar: columna "usuario" en la plantilla .xlsx manda por fila.

La plantilla descargable lleva una columna final `usuario` con el username del
dueño de la información. La celda manda por fila: si está vacía o la columna
no existe (archivos viejos), el responsable es el usuario que ejecuta la
importación. Username desconocido o inactivo deja la fila inválida en la
vista previa. El registro (userId) queda en el responsable; audit_logs en el
actor real. El selector visible del form ya no existe: manda el Excel.
"""

import json
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.models.audit_log import AuditLog
from app.models.devolucion_folio import DevolucionFolio
from app.models.entrega_folio import EntregaFolio
from app.models.recepcion_folio import RecepcionFolio
from tests.conftest import make_folio, make_user

HEADERS_RECEPCION = ['fecha', 'anioCert', 'tipoCert', 'folioInicial', 'folioFinal', 'usuario']
HEADERS_RECEPCION_SIN_USUARIO = HEADERS_RECEPCION[:-1]
HEADERS_FOLIOS = ['folio', 'tipoCert', 'anioCert', 'usuario']
Y = 2026


def make_xlsx(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


@pytest.fixture()
def otro_usuario(app):
    """Segundo usuario activo: el dueño de la información importada."""
    return make_user('otro', 'operador')


# --- la plantilla descargable trae la columna usuario ---


@pytest.mark.parametrize('url', [
    '/recepcion/plantilla',
    '/entregas/plantilla',
    '/devoluciones/plantilla',
])
def test_plantilla_incluye_columna_usuario(admin_client, url):
    resp = admin_client.get(url)
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    ws = wb.active
    encabezados = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert encabezados[-1] == 'usuario'
    ejemplo = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert ejemplo[-1] == 'adminuser', 'la fila de ejemplo lleva el username actual'


# --- el selector visible desaparece: manda el Excel ---


@pytest.mark.parametrize('url', [
    '/recepcion/importar',
    '/entregas/importar',
    '/devoluciones/importar',
])
def test_importar_sin_selector_de_usuario_visible(admin_client, url):
    text = admin_client.get(url).get_data(as_text=True)
    assert 'id="usuarioId"' not in text
    assert 'Usuario responsable' not in text


# --- recepción: la celda manda por fila ---


def test_recepcion_columna_usuario_atribuye_a_otro(admin_client, tipo, admin, otro_usuario):
    resp = admin_client.post(
        '/recepcion/importar',
        data={'archivo': (
            make_xlsx(HEADERS_RECEPCION,
                      [['2026-10-01', Y, tipo.name, 1, 5, 'otro']]),
            'i.xlsx')},
        content_type='multipart/form-data',
    )
    text = resp.get_data(as_text=True)
    assert 'usuario_id' in text, 'el payload lleva el usuario por fila'

    payload = json.dumps([{
        'fecha': '2026-10-01', 'anioCert': Y, 'tipoCert_id': tipo.id,
        'folioInicial': 1, 'folioFinal': 5, 'usuario_id': otro_usuario.id,
    }])
    admin_client.post('/recepcion/importar/confirmar',
                     data={'payload': payload}, follow_redirects=True)

    rec = RecepcionFolio.query.one()
    assert rec.userId == otro_usuario.id, 'el registro queda en el responsable'
    log = AuditLog.query.filter_by(tabla='recepcion_folios', accion='INSERT').one()
    assert log.user_id == admin.id, 'la auditoría queda en el actor real'


def test_recepcion_usuario_desconocido_marca_fila(admin_client, tipo):
    resp = admin_client.post(
        '/recepcion/importar',
        data={'archivo': (
            make_xlsx(HEADERS_RECEPCION,
                      [['2026-10-01', Y, tipo.name, 1, 5, 'fantasma']]),
            'i.xlsx')},
        content_type='multipart/form-data',
    )
    text = resp.get_data(as_text=True)
    assert 'Usuario desconocido: fantasma.' in text
    assert 'Ninguna fila del archivo es válida' in text


def test_recepcion_usuario_inactivo_marca_fila(admin_client, tipo, app):
    baja = make_user('baja', 'operador')
    baja.activo = False
    from app.extensions import db
    db.session.commit()

    resp = admin_client.post(
        '/recepcion/importar',
        data={'archivo': (
            make_xlsx(HEADERS_RECEPCION,
                      [['2026-10-01', Y, tipo.name, 1, 5, 'baja']]),
            'i.xlsx')},
        content_type='multipart/form-data',
    )
    assert 'Usuario inactivo: baja.' in resp.get_data(as_text=True)


def test_recepcion_sin_columna_atribuye_al_actual(admin_client, tipo, admin):
    resp = admin_client.post(
        '/recepcion/importar',
        data={'archivo': (
            make_xlsx(HEADERS_RECEPCION_SIN_USUARIO,
                      [['2026-10-01', Y, tipo.name, 1, 5]]),
            'i.xlsx')},
        content_type='multipart/form-data',
    )
    assert 'Válida' in resp.get_data(as_text=True), 'archivos viejos siguen validos'

    payload = json.dumps([{
        'fecha': '2026-10-01', 'anioCert': Y, 'tipoCert_id': tipo.id,
        'folioInicial': 1, 'folioFinal': 5,
    }])
    admin_client.post('/recepcion/importar/confirmar',
                      data={'payload': payload}, follow_redirects=True)
    assert RecepcionFolio.query.one().userId == admin.id


def test_recepcion_celda_vacia_atribuye_al_actual(admin_client, tipo, admin, otro_usuario):
    resp = admin_client.post(
        '/recepcion/importar',
        data={'archivo': (
            make_xlsx(HEADERS_RECEPCION,
                      [['2026-10-01', Y, tipo.name, 1, 5, '']]),
            'i.xlsx')},
        content_type='multipart/form-data',
    )
    text = resp.get_data(as_text=True)
    assert f'&#34;usuario_id&#34;: {admin.id}' in text
    assert 'Válida' in text


def test_recepcion_preview_muestra_responsable_por_fila(admin_client, tipo, otro_usuario):
    resp = admin_client.post(
        '/recepcion/importar',
        data={'archivo': (
            make_xlsx(HEADERS_RECEPCION,
                      [['2026-10-01', Y, tipo.name, 1, 5, 'otro']]),
            'i.xlsx')},
        content_type='multipart/form-data',
    )
    text = resp.get_data(as_text=True)
    assert 'otro' in text
    assert 'a nombre de' in text


# --- entregas: atribución mixta por fila y agrupación en el confirm ---


def test_entregas_preview_lleva_usuario_por_folio(admin_client, tipo, centro, operador, otro_usuario):
    make_folio(tipo, anio=Y, numero=1000)
    make_folio(tipo, anio=Y, numero=1001)
    resp = admin_client.post(
        '/entregas/importar',
        data={
            'archivo': (make_xlsx(HEADERS_FOLIOS, [
                [1000, tipo.name, Y, 'otro'],
                [1001, tipo.name, Y, 'operador'],
            ]), 'i.xlsx'),
            'centroId': str(centro.id),
            'fechaEntrega': '2026-10-01',
        },
        content_type='multipart/form-data',
    )
    text = resp.get_data(as_text=True)
    assert 'usuarioId' in text, 'el payload lleva usuarioId por folio'


def test_entregas_confirm_agrupa_por_usuario(admin_client, tipo, centro, admin, operador, otro_usuario):
    f1 = make_folio(tipo, anio=Y, numero=1000)
    f2 = make_folio(tipo, anio=Y, numero=1001)
    admin_client.post('/entregas/importar/confirmar', data={
        'payload': json.dumps([
            {'id': f1.id, 'usuarioId': otro_usuario.id},
            {'id': f2.id, 'usuarioId': operador.id},
        ]),
        'centroId': str(centro.id),
        'fecha': '2026-10-01',
    }, follow_redirects=True)

    entregas = EntregaFolio.query.order_by(EntregaFolio.folio_id).all()
    assert len(entregas) == 2
    assert {e.userId for e in entregas} == {otro_usuario.id, operador.id}
    logs = AuditLog.query.filter_by(tabla='entrega_folios', accion='INSERT').all()
    assert len(logs) == 2
    assert all(l.user_id == admin.id for l in logs), 'audit siempre al actor'


def test_entregas_usuario_desconocido_marca_fila(admin_client, tipo, centro):
    make_folio(tipo, anio=Y, numero=1000)
    resp = admin_client.post(
        '/entregas/importar',
        data={
            'archivo': (make_xlsx(HEADERS_FOLIOS, [[1000, tipo.name, Y, 'fantasma']]),
                        'i.xlsx'),
            'centroId': str(centro.id),
            'fechaEntrega': '2026-10-01',
        },
        content_type='multipart/form-data',
    )
    assert 'Usuario desconocido: fantasma.' in resp.get_data(as_text=True)


# --- devoluciones ---


def test_devoluciones_columna_usuario_atribuye_a_otro(admin_client, folio_entregado, centro,
                                                      admin, otro_usuario):
    resp = admin_client.post(
        '/devoluciones/importar',
        data={
            'archivo': (make_xlsx(HEADERS_FOLIOS,
                                  [[folio_entregado.folio,
                                    'Defunciones', Y, 'otro']]),
                        'i.xlsx'),
            'centroId': str(centro.id),
            'fechaDevolucion': '2026-10-02',
        },
        content_type='multipart/form-data',
    )
    assert 'usuarioId' in resp.get_data(as_text=True)

    admin_client.post('/devoluciones/importar/confirmar', data={
        'payload': json.dumps([{'id': folio_entregado.id,
                                'usuarioId': otro_usuario.id}]),
        'centroId': str(centro.id),
        'fecha': '2026-10-02',
    }, follow_redirects=True)

    d = DevolucionFolio.query.one()
    assert d.userId == otro_usuario.id
    log = AuditLog.query.filter_by(tabla='devolucion_folios', accion='INSERT').one()
    assert log.user_id == admin.id


def test_devoluciones_usuario_desconocido_marca_fila(admin_client, folio_entregado, centro):
    resp = admin_client.post(
        '/devoluciones/importar',
        data={
            'archivo': (make_xlsx(HEADERS_FOLIOS,
                                  [[folio_entregado.folio,
                                    'Defunciones', Y, 'fantasma']]),
                        'i.xlsx'),
            'centroId': str(centro.id),
            'fechaDevolucion': '2026-10-02',
        },
        content_type='multipart/form-data',
    )
    assert 'Usuario desconocido: fantasma.' in resp.get_data(as_text=True)


# --- defensas en el confirm (payload es editable) ---


def test_confirm_payload_enteros_legado_usa_al_actual(admin_client, tipo, centro, admin):
    folio = make_folio(tipo, anio=Y, numero=1000)
    admin_client.post('/entregas/importar/confirmar', data={
        'payload': json.dumps([folio.id]),
        'centroId': str(centro.id),
        'fecha': '2026-10-01',
    }, follow_redirects=True)
    assert EntregaFolio.query.one().userId == admin.id


def test_confirm_usuario_id_invalido_usa_al_actual(admin_client, tipo, centro, admin):
    folio = make_folio(tipo, anio=Y, numero=1000)
    admin_client.post('/entregas/importar/confirmar', data={
        'payload': json.dumps([{'id': folio.id, 'usuarioId': 999999}]),
        'centroId': str(centro.id),
        'fecha': '2026-10-01',
    }, follow_redirects=True)
    assert EntregaFolio.query.one().userId == admin.id
