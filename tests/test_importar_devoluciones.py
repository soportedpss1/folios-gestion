"""Importar devoluciones de folios desde .xlsx (vista previa + confirmación).

La fecha va por fila en el Excel (columna obligatoria `fecha`): no hay campo
de fecha en el formulario ni hidden en la confirmación. Cada movimiento
guarda la fecha de su fila.

El centro de salud TAMBIÉN viene por fila: la columna `centro` lleva el
nombre exacto del centro en la DB. Solo pasa el folio si ese centro de la
fila lo tiene entregado; un archivo puede mezclar centros distintos.

TODAS las columnas son obligatorias (incluida la final `usuario`, con el
username del responsable): un archivo sin ella se rechaza y una celda
vacía deja la fila inválida en la vista previa.
"""

import json
from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models.centro import Centro
from app.models.devolucion_folio import DevolucionFolio
from app.models.entrega_folio import EntregaFolio
from app.models.folio import Folio
from tests.conftest import make_folio

HEADERS = ['folio', 'tipoCert', 'anioCert', 'fecha', 'centro', 'usuario']
Y = 2026


def make_xlsx(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def post_importar(client, rows, nombre='import.xlsx'):
    return client.post(
        '/devoluciones/importar',
        data={'archivo': (make_xlsx(HEADERS, rows), nombre)},
        content_type='multipart/form-data',
    )


def confirmar(client, items):
    """items: lista de {'id', 'fecha', 'centroId', ['usuarioId']}."""
    return client.post(
        '/devoluciones/importar/confirmar',
        data={'payload': json.dumps(items)},
        follow_redirects=True,
    )


def test_importar_requiere_login(client):
    assert client.get('/devoluciones/importar').status_code == 302
    assert client.post('/devoluciones/importar/confirmar').status_code == 302


def test_plantilla_es_xlsx_con_headers(admin_client, centro):
    resp = admin_client.get('/devoluciones/plantilla')
    assert resp.status_code == 200
    assert 'spreadsheetml' in resp.content_type
    wb = load_workbook(BytesIO(resp.data))
    assert [c.value for c in wb.active[1]] == HEADERS, \
        'la plantilla lleva todas las columnas obligatorias (usuario al final)'
    ejemplo = [c.value for c in wb.active[2]]
    assert ejemplo[4] == centro.name, 'la fila de ejemplo lleva un centro real'


def test_formulario_no_tiene_campo_fecha(admin_client, centro):
    text = admin_client.get('/devoluciones/importar').get_data(as_text=True)
    assert 'id="fechaDevolucion"' not in text, 'la fecha se toma del Excel'


def test_formulario_no_tiene_selector_de_centro(admin_client):
    text = admin_client.get('/devoluciones/importar').get_data(as_text=True)
    assert 'id="centroId"' not in text, 'el centro se toma del Excel'


def test_preview_resuelve_entregado_al_centro(admin_client, folio_entregado,
                                              centro, tipo):
    resp = post_importar(admin_client, [
        [folio_entregado.folio, tipo.name, Y, '02/10/2026', centro.name,
         'adminuser']])
    text = resp.get_data(as_text=True)
    assert '<strong>1</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'Confirmar importación' in text
    assert 'Todas las filas del archivo son válidas' in text, \
        'sin inválidas no se lista ninguna fila (la fecha por fila se verifica en la lista de inválidas)'
    assert '<input type="hidden" name="fecha"' not in text, 'sin hidden de fecha'
    assert 'id="centroId"' not in text, 'sin hidden/select de centro'
    assert DevolucionFolio.query.count() == 0, 'la vista previa no debe crear nada'


def test_preview_rechaza_folio_de_otro_centro(admin_client, folio_entregado, tipo):
    otro = Centro(name='Otro Centro', telefono='555', direccion='Calle 1', activo=True)
    db.session.add(otro)
    db.session.commit()

    resp = post_importar(admin_client, [
        [folio_entregado.folio, tipo.name, Y, '2026-10-02', otro.name,
         'adminuser']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'no lo tiene entregado' in text


def test_preview_rechaza_folio_no_entregado(admin_client, tipo, centro):
    make_folio(tipo, anio=Y, numero=1000, estado='disponible')
    resp = post_importar(admin_client,
                         [[1000, tipo.name, Y, '2026-10-02', centro.name,
                           'adminuser']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'no está entregado' in text


def test_preview_lista_solo_filas_invalidas(admin_client, folio_entregado,
                                            centro, tipo):
    resp = post_importar(admin_client, [
        [folio_entregado.folio, tipo.name, Y, '02/10/2026', centro.name,
         'adminuser'],
        [9999, tipo.name, Y, '02/10/2026', centro.name, 'adminuser'],
    ])
    text = resp.get_data(as_text=True)
    assert '<strong>1</strong> folio(s) válido(s) de <strong>2</strong>' in text, \
        'el resumen sigue contando válidas y total'
    assert 'El folio no existe.' in text, 'la fila inválida se lista con su error'
    assert '02/10/2026' in text, 'la fila inválida muestra su fecha por fila'
    assert f'Folio {folio_entregado.folio}' not in text, \
        'la lista de validación no debe mostrar filas válidas'


def test_preview_sin_filas_invalidas_muestra_mensaje(admin_client,
                                                     folio_entregado, centro,
                                                     tipo):
    resp = post_importar(admin_client, [
        [folio_entregado.folio, tipo.name, Y, '02/10/2026', centro.name,
         'adminuser']])
    text = resp.get_data(as_text=True)
    assert 'Todas las filas del archivo son válidas' in text, \
        'sin inválidas se muestra el mensaje en lugar de la tabla'
    assert f'Folio {folio_entregado.folio}' not in text, \
        'no debe haber filas en la lista'
    assert 'Confirmar importación' in text, 'el resumen y la confirmación siguen'


def test_preview_centro_desconocido_marca_fila(admin_client, folio_entregado, tipo):
    resp = post_importar(admin_client, [
        [folio_entregado.folio, tipo.name, Y, '2026-10-02', 'No Existe',
         'adminuser']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'Centro de salud desconocido: No Existe.' in text


def test_preview_centro_inactivo_marca_fila(admin_client, folio_entregado, tipo):
    baja = Centro(name='Centro Baja', telefono='555', direccion='Calle 2',
                  activo=False)
    db.session.add(baja)
    db.session.commit()

    resp = post_importar(admin_client, [
        [folio_entregado.folio, tipo.name, Y, '2026-10-02', 'Centro Baja',
         'adminuser']])
    assert 'Centro de salud inactivo: Centro Baja.' in \
        resp.get_data(as_text=True)


def test_preview_admite_centros_mezclados(admin_client, folio_entregado,
                                          centro, tipo, operador):
    otro = Centro(name='Otro Centro', telefono='555', direccion='Calle 1',
                  activo=True)
    db.session.add(otro)
    db.session.commit()
    segundo = make_folio(tipo, anio=Y, numero=1001, estado='entregado')
    db.session.add(EntregaFolio(folio_id=segundo.id,
                                fechaEntrega=date(2026, 9, 1),
                                centroId=otro.id, userId=operador.id))
    db.session.commit()

    resp = post_importar(admin_client, [
        [folio_entregado.folio, tipo.name, Y, '2026-10-02', centro.name,
         'adminuser'],
        [segundo.folio, tipo.name, Y, '2026-10-02', otro.name, 'adminuser'],
    ])
    text = resp.get_data(as_text=True)
    assert '<strong>2</strong> folio(s) válido(s) de <strong>2</strong>' in text
    assert f'&#34;centroId&#34;: {centro.id}' in text, 'payload con centro por fila'
    assert f'&#34;centroId&#34;: {otro.id}' in text


def test_preview_rechaza_archivo_sin_columna_fecha(admin_client, folio_entregado,
                                                   centro, tipo):
    resp = admin_client.post('/devoluciones/importar', data={
        'archivo': (make_xlsx(['folio', 'tipoCert', 'anioCert', 'centro'],
                              [[folio_entregado.folio, tipo.name, Y,
                                centro.name]]),
                    'sin-fecha.xlsx'),
    }, content_type='multipart/form-data')
    assert 'Encabezados incorrectos' in resp.get_data(as_text=True)


def test_preview_rechaza_archivo_sin_columna_centro(admin_client, folio_entregado,
                                                    tipo):
    resp = admin_client.post('/devoluciones/importar', data={
        'archivo': (make_xlsx(['folio', 'tipoCert', 'anioCert', 'fecha'],
                              [[folio_entregado.folio, tipo.name, Y,
                                '2026-10-02']]),
                    'sin-centro.xlsx'),
    }, content_type='multipart/form-data')
    assert 'Encabezados incorrectos' in resp.get_data(as_text=True)


def test_preview_rechaza_archivo_sin_columna_usuario(admin_client, folio_entregado,
                                                     centro, tipo):
    """Todas las columnas son obligatorias: sin `usuario` no hay importación."""
    resp = admin_client.post('/devoluciones/importar', data={
        'archivo': (make_xlsx(HEADERS[:-1],
                              [[folio_entregado.folio, tipo.name, Y,
                                '2026-10-02', centro.name]]),
                    'sin-usuario.xlsx'),
    }, content_type='multipart/form-data')
    assert 'Encabezados incorrectos' in resp.get_data(as_text=True)


def test_preview_celda_usuario_vacia_marca_fila(admin_client, folio_entregado,
                                                centro, tipo):
    """La celda también es obligatoria: vacía → fila inválida."""
    resp = post_importar(admin_client, [
        [folio_entregado.folio, tipo.name, Y, '2026-10-02', centro.name, '']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'Usuario vacío.' in text


def test_preview_marca_fecha_invalida(admin_client, folio_entregado, centro, tipo):
    resp = post_importar(admin_client, [
        [folio_entregado.folio, tipo.name, Y, 'mañana', centro.name,
         'adminuser']])
    text = resp.get_data(as_text=True)
    assert '<strong>0</strong> folio(s) válido(s) de <strong>1</strong>' in text
    assert 'Fecha inválida' in text


def test_confirm_registra_devolucion(admin_client, folio_entregado, centro, tipo):
    resp = confirmar(admin_client, [
        {'id': folio_entregado.id, 'fecha': '2026-10-02',
         'centroId': centro.id},
    ])

    assert resp.status_code == 200
    assert DevolucionFolio.query.count() == 1
    assert db.session.get(Folio, folio_entregado.id).estado == 'devuelto'
    devolucion = DevolucionFolio.query.one()
    assert devolucion.fechaDevolucion == date(2026, 10, 2)
    assert devolucion.centroId == centro.id
    assert '1 devolución registrada' in resp.get_data(as_text=True)


def test_confirm_respeta_centro_distinto_por_fila(admin_client, folio_entregado,
                                                  centro, tipo, operador):
    otro = Centro(name='Otro Centro', telefono='555', direccion='Calle 1',
                  activo=True)
    db.session.add(otro)
    db.session.commit()
    segundo = make_folio(tipo, anio=Y, numero=1001, estado='entregado')
    db.session.add(EntregaFolio(folio_id=segundo.id,
                                fechaEntrega=date(2026, 9, 1),
                                centroId=otro.id, userId=operador.id))
    db.session.commit()

    resp = confirmar(admin_client, [
        {'id': folio_entregado.id, 'fecha': '2026-10-02', 'centroId': centro.id},
        {'id': segundo.id, 'fecha': '2026-10-02', 'centroId': otro.id},
    ])

    assert resp.status_code == 200
    assert DevolucionFolio.query.count() == 2
    assert {d.centroId for d in DevolucionFolio.query.all()} == \
        {centro.id, otro.id}, 'cada folio se devuelve al centro de su fila'


def test_confirm_respeta_fecha_distinta_por_fila(admin_client, folio_entregado,
                                                 centro, tipo, operador):
    otro = make_folio(tipo, anio=Y, numero=1001, estado='entregado')
    db.session.add(EntregaFolio(folio_id=otro.id, fechaEntrega=date(2026, 9, 1),
                                centroId=centro.id, userId=operador.id))
    db.session.commit()

    resp = confirmar(admin_client, [
        {'id': folio_entregado.id, 'fecha': '2026-10-02', 'centroId': centro.id},
        {'id': otro.id, 'fecha': '07/10/2026', 'centroId': centro.id},
    ])

    assert resp.status_code == 200
    assert DevolucionFolio.query.count() == 2
    assert {d.fechaDevolucion for d in DevolucionFolio.query.all()} == \
        {date(2026, 10, 2), date(2026, 10, 7)}


def test_confirm_revalida_centro(admin_client, folio_entregado, tipo):
    """Folio entregado a otro centro → confirm no lo devuelve."""
    otro = Centro(name='Otro Centro', telefono='555', direccion='Calle 1', activo=True)
    db.session.add(otro)
    db.session.commit()

    resp = confirmar(admin_client, [
        {'id': folio_entregado.id, 'fecha': '2026-10-02', 'centroId': otro.id},
    ])

    assert DevolucionFolio.query.count() == 0
    assert 'Ningún folio sigue entregado' in resp.get_data(as_text=True)
    assert db.session.get(Folio, folio_entregado.id).estado == 'entregado'


def test_confirm_payload_invalido_no_crea_nada(admin_client, folio_entregado,
                                               centro):
    resp = admin_client.post('/devoluciones/importar/confirmar', data={
        'payload': 'no-es-json',
    }, follow_redirects=True)

    assert 'inválidos' in resp.get_data(as_text=True)
    assert DevolucionFolio.query.count() == 0


def test_confirm_rechaza_fecha_invalida_en_payload(admin_client, folio_entregado,
                                                   centro):
    """El hidden del confirm es editable: fecha faltante o basura → rechazo."""
    resp = confirmar(admin_client, [
        {'id': folio_entregado.id, 'centroId': centro.id},
    ])

    assert 'inválidos' in resp.get_data(as_text=True)
    assert DevolucionFolio.query.count() == 0


def test_confirm_rechaza_centro_ausente_en_payload(admin_client, folio_entregado):
    """Sin `centroId` por fila el payload completo es inválido."""
    resp = confirmar(admin_client, [
        {'id': folio_entregado.id, 'fecha': '2026-10-02'},
    ])

    assert 'inválidos' in resp.get_data(as_text=True)
    assert DevolucionFolio.query.count() == 0


def test_confirm_rechaza_centro_desconocido(admin_client, folio_entregado, centro):
    """El hidden es editable: centro inexistente → nada se registra."""
    resp = confirmar(admin_client, [
        {'id': folio_entregado.id, 'fecha': '2026-10-02', 'centroId': 999999},
    ])

    assert 'Centro de salud inválido.' in resp.get_data(as_text=True)
    assert DevolucionFolio.query.count() == 0
    assert db.session.get(Folio, folio_entregado.id).estado == 'entregado'
