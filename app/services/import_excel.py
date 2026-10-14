"""Lectura de plantillas .xlsx para importación masiva.

Recepción, entregas y devoluciones leen el mismo tipo de archivo: una hoja con
encabezados fijos en la fila 1 y una fila de datos por registro. Aquí vive lo
genérico —lectura, coerción de celdas, parseo del confirm y descarga de
plantilla—; la validación de negocio queda en cada blueprint.
"""

import json
import zipfile
from datetime import date, datetime
from io import BytesIO

from flask import make_response
from openpyxl import Workbook, load_workbook

from app.services.excel import neutralizar_formulas

# Tope de filas por archivo de folios (entregas/devoluciones). El de rangos de
# recepción lo define su blueprint porque allí cada fila es un rango completo.
MAX_FOLIOS_POR_ARCHIVO = 5000

# MAX_CONTENT_LENGTH acota el .xlsx COMPRIMIDO; openpyxl carga en memoria el
# XML descomprimido (sharedStrings entero) → un zip de unos pocos MB puede
# descomprimirse en gigas y tumbar al worker (OOM).
MAX_XLSX_DECOMPRESION = 200 * 1024 * 1024

# Columna final de las plantillas: username del dueño de la fila.
# Obligatoria en TODAS las importaciones (columna y celda por fila).
COLUMNA_USUARIO = 'usuario'


class ImportExcelError(Exception):
    """Archivo ilegible o con formato incorrecto. El mensaje ya está en español."""


def _validar_tamano_descomprimido(file_storage):
    """Rechaza .xlsx cuyo contenido descomprimido supera MAX_XLSX_DECOMPRESION."""
    try:
        with zipfile.ZipFile(file_storage.stream) as zf:
            total = sum(info.file_size for info in zf.infolist())
    except zipfile.BadZipFile:
        return  # load_workbook dará el mensaje genérico de "no es un Excel"
    finally:
        file_storage.stream.seek(0)
    if total > MAX_XLSX_DECOMPRESION:
        mb = MAX_XLSX_DECOMPRESION // (1024 * 1024)
        raise ImportExcelError(
            f'El archivo ocupa demasiado al descomprimirse (máximo {mb} MB).')


def leer_filas(file_storage, headers, max_filas):
    """Devuelve las filas de datos como dicts con llave = encabezado.

    `headers` es la lista completa (en orden) que debe tener la fila 1:
    TODAS las columnas son obligatorias y no puede sobrar ninguna. Lanza
    ImportExcelError si los encabezados no coinciden, si la hoja no tiene
    datos o si supera `max_filas` filas. Las filas totalmente vacías se
    ignoran.
    """
    try:
        _validar_tamano_descomprimido(file_storage)
        wb = load_workbook(file_storage.stream, read_only=True, data_only=True)
    except ImportExcelError:
        raise
    except Exception:
        raise ImportExcelError('El archivo no es un Excel válido (.xlsx).')

    try:
        ws = wb.active
        esperados = [h.strip().lower() for h in headers]
        filas = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                actual = [('' if c is None else str(c).strip().lower()) for c in row]
                while actual and actual[-1] == '':
                    actual.pop()
                if actual != esperados:
                    raise ImportExcelError(
                        'Encabezados incorrectos. Se esperaba exactamente: '
                        + ', '.join(headers) + '.'
                    )
                continue
            if all(c is None or str(c).strip() == '' for c in row):
                continue
            if len(filas) >= max_filas:
                raise ImportExcelError(
                    f'El archivo supera el máximo de {max_filas} filas.'
                )
            filas.append({h: v for h, v in zip(headers, row)})
    finally:
        wb.close()

    if not filas:
        raise ImportExcelError('El archivo no tiene filas de datos.')
    return filas


def a_entero(valor):
    """Celda → entero, o None si no representa un número entero."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, int):
        return valor
    if isinstance(valor, float):
        return int(valor) if valor.is_integer() else None
    try:
        return int(str(valor).strip())
    except ValueError:
        return None


def a_fecha(valor):
    """Celda → date (datetime, date, 'dd/mm/aaaa' o 'aaaa-mm-dd'), o None."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if valor is None:
        return None
    texto = str(valor).strip()
    for formato in ('%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def parsear_confirmacion(payload_raw):
    """Hidden input del confirm → (items, error).

    `items` es una lista de (folio_id, usuario_id|None, fecha, centro_id):
    el payload de entregas/devoluciones lleva responsable, fecha y centro por
    fila (la fecha viene del Excel, columna obligatoria `fecha`; el centro
    también, columna obligatoria `centro` — ya no hay selects en el
    formulario). Aquí se sanea porque el hidden es editable; el estado de cada
    folio y el centro se reconsultan después, en el llamador.
    """
    invalido = 'Datos de importación inválidos. Repita la operación.'
    try:
        datos = json.loads(payload_raw or '')
    except (TypeError, ValueError):
        return None, invalido
    if not isinstance(datos, list) or not datos \
            or len(datos) > MAX_FOLIOS_POR_ARCHIVO:
        return None, invalido

    items = []
    for dato in datos:
        if not isinstance(dato, dict):
            return None, invalido
        fid, uid = dato.get('id'), dato.get('usuarioId')
        if not isinstance(fid, int) or isinstance(fid, bool) or fid <= 0:
            return None, invalido
        if uid is not None and (not isinstance(uid, int) or isinstance(uid, bool)
                                 or uid <= 0):
            return None, invalido
        fecha = a_fecha(dato.get('fecha'))
        if fecha is None:
            return None, invalido
        centro_id = a_entero(dato.get('centroId'))
        if centro_id is None or centro_id <= 0:
            return None, invalido
        items.append((fid, uid, fecha, centro_id))
    return items, None


def respuesta_plantilla(nombre_archivo, headers, fila_ejemplo):
    """Response .xlsx con encabezados formateados + fila de ejemplo."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Plantilla'

    from openpyxl.styles import Alignment, Font, PatternFill
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    for col, valor in enumerate(fila_ejemplo, 1):
        ws.cell(row=2, column=col, value=valor)
    for col in ws.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 2

    output = BytesIO()
    neutralizar_formulas(wb)
    wb.save(output)
    output.seek(0)

    response = make_response(output.getvalue())
    response.headers['Content-Type'] = (
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response.headers['Content-Disposition'] = f'attachment; filename={nombre_archivo}'
    return response
