import json
from datetime import date, datetime
from io import BytesIO

from flask import (Blueprint, render_template, redirect, url_for, flash, request,
                   make_response)
from flask_login import login_required, current_user
from markupsafe import escape
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import selectinload
from app.extensions import db
from app.models.recepcion_folio import RecepcionFolio
from app.models.folio import Folio
from app.models.tipo_certificado import TipoCertificado
from app.recepcion.forms import (
    ImportarRecepcionForm, RecepcionForm, detectar_solape,
    MAX_FOLIOS_POR_RECEPCION,
)
from app.services import import_excel
from app.services.audit import log_audit
from app.services.excel import neutralizar_formulas
from app.services.mensajes import msg_solape, INTENTE_DE_NUEVO
from app.services.pdf import encabezado_pdf
from app.services.queries import resolver_usuario_nombre, usuario_activo
from app.utils import plural
from app.decorators import permiso_requerido

recepcion_bp = Blueprint('recepcion', __name__, url_prefix='/recepcion')

# Tope de rangos por archivo de importación: cada fila crea hasta
# MAX_FOLIOS_POR_RECEPCION folios individuales.
MAX_RANGOS_POR_ARCHIVO = 500
# Columna final opcional: username del dueño del rango (fila = usuario actual
# si falta o está vacía).
HEADERS_RECEPCION = ['fecha', 'anioCert', 'tipoCert', 'folioInicial', 'folioFinal',
                     import_excel.COLUMNA_USUARIO]
MAX_PREVIA_MOSTRAR = 200


class SolapeError(Exception):
    """Rango nuevo en conflicto con una recepción existente."""


def _next_rango_id():
    """MAX(rangoId)+1 con bloqueo de fila para serializar recepciones concurrentes."""
    last = RecepcionFolio.query.order_by(RecepcionFolio.rangoId.desc()) \
        .with_for_update().first()
    return (last.rangoId + 1) if last else 1


def registrar_rango(*, fecha, anioCert, tipoCert_id, folioInicial, folioFinal,
                    user_id, audit_user_id=None):
    """Crea la RecepcionFolio y sus folios individuales, sin commit.

    `user_id` es el responsable del registro; `audit_user_id` quién ejecuta la
    operación y queda en audit_logs (por defecto, el mismo).

    Corre bajo el lock de _next_rango_id y rechequea el solape ahí: la
    validación del form corre antes del lock, así que dos POST concurrentes
    ambos pasan allí. Lanza SolapeError (con rollback hecho) o IntegrityError;
    el llamador decide el commit.
    """
    new_rango_id = _next_rango_id()

    solape = detectar_solape(anioCert, tipoCert_id, folioInicial, folioFinal)
    if solape:
        db.session.rollback()
        raise SolapeError(msg_solape(folioInicial, folioFinal,
                                     solape.folioInicial, solape.folioFinal))

    recepcion = RecepcionFolio(
        fecha=fecha,
        anioCert=anioCert,
        tipoCert_id=tipoCert_id,
        folioInicial=folioInicial,
        folioFinal=folioFinal,
        rangoId=new_rango_id,
        userId=user_id,
    )
    db.session.add(recepcion)
    db.session.flush()

    # bulk insert: un solo INSERT multi-fila en vez de uno por folio
    db.session.bulk_insert_mappings(Folio, [
        {
            'rangoId': new_rango_id,
            'anioCert': anioCert,
            'tipoCert_id': tipoCert_id,
            'folio': num,
            'digitado': False,
            'escaneado': False,
            'nulo': False,
            'estado': 'disponible',
        }
        for num in range(folioInicial, folioFinal + 1)
    ])

    log_audit(audit_user_id or user_id, 'INSERT', 'recepcion_folios', recepcion.id, None,
              recepcion.to_dict())
    return folioFinal - folioInicial + 1


@recepcion_bp.route('/')
@login_required
@permiso_requerido('recepcion.ver')
def index():
    page = request.args.get('page', 1, type=int)
    pagination = RecepcionFolio.query.options(
        selectinload(RecepcionFolio.tipo_cert),
        selectinload(RecepcionFolio.usuario),
    ).order_by(RecepcionFolio.created_at.desc()).paginate(page=page, per_page=25, error_out=False)
    return render_template('recepcion/index.html', recepciones=pagination.items, pagination=pagination)


def _recepciones_para_exportar():
    return RecepcionFolio.query.options(
        selectinload(RecepcionFolio.tipo_cert),
        selectinload(RecepcionFolio.usuario),
    ).order_by(RecepcionFolio.id).all()


@recepcion_bp.route('/export/excel')
@login_required
@permiso_requerido('recepcion.ver')
def export_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    recepciones = _recepciones_para_exportar()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Recepciones'

    headers = ['#', 'Fecha', 'Año Cert.', 'Tipo', 'Folio Inicial', 'Folio Final',
               'Total Folios', 'Registrado Por']
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
    header_align = Alignment(horizontal='center')
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for row, r in enumerate(recepciones, 2):
        ws.cell(row=row, column=1, value=r.id)
        ws.cell(row=row, column=2, value=r.fecha.strftime('%d/%m/%Y') if r.fecha else '')
        ws.cell(row=row, column=3, value=r.anioCert)
        ws.cell(row=row, column=4, value=r.tipo_cert.name if r.tipo_cert else '')
        ws.cell(row=row, column=5, value=r.folioInicial)
        ws.cell(row=row, column=6, value=r.folioFinal)
        ws.cell(row=row, column=7, value=r.folioFinal - r.folioInicial + 1)
        ws.cell(row=row, column=8, value=r.usuario.username if r.usuario else '')

    for col in ws.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2

    output = BytesIO()
    neutralizar_formulas(wb)  # nunca escribir '=...' como fórmula (inyección al abrir)
    wb.save(output)
    output.seek(0)

    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = 'attachment; filename=recepciones.xlsx'
    return response


@recepcion_bp.route('/export/pdf')
@login_required
@permiso_requerido('recepcion.ver')
def export_pdf():
    from weasyprint import HTML

    recepciones = _recepciones_para_exportar()

    encabezado = encabezado_pdf('Recepcion de Folios')
    html_content = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <style>
            @page {{ size: landscape; margin: 1.5cm; }}
            body {{ font-family: Arial, sans-serif; margin: 20px; font-size: 10px; }}
            h1 {{ color: #0d6efd; text-align: center; font-size: 16px; }}
            h2 {{ color: #333; text-align: center; font-size: 13px; margin-top: 5px; }}
            h3 {{ color: #666; text-align: center; font-size: 11px; margin-top: 5px; }}
            table {{ width: 100%; border-collapse: collapse; margin-top: 15px; font-size: 9px; }}
            th, td {{ border: 1px solid #ddd; padding: 5px 6px; text-align: left; }}
            th {{ background-color: #0d6efd; color: white; font-size: 9px; }}
            tr:nth-child(even) {{ background-color: #f2f2f2; }}
            .header {{ text-align: center; margin-bottom: 20px; }}
            .date {{ color: #666; font-size: 10px; text-align: right; margin-bottom: 10px; }}
        </style>
    </head>
    <body>
        {encabezado}
        <p class="date">Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}</p>
        <table>
            <thead>
                <tr>
                    <th>#</th>
                    <th>Fecha</th>
                    <th>Ano Cert.</th>
                    <th>Tipo</th>
                    <th>Folio Inicial</th>
                    <th>Folio Final</th>
                    <th>Total Folios</th>
                    <th>Registrado Por</th>
                </tr>
            </thead>
            <tbody>
    """

    filas = []
    for r in recepciones:
        filas.append(f"""
                <tr>
                    <td>{r.id}</td>
                    <td>{r.fecha.strftime('%d/%m/%Y') if r.fecha else '-'}</td>
                    <td>{r.anioCert}</td>
                    <td>{escape(r.tipo_cert.name) if r.tipo_cert else ''}</td>
                    <td>{r.folioInicial}</td>
                    <td>{r.folioFinal}</td>
                    <td>{r.folioFinal - r.folioInicial + 1}</td>
                    <td>{escape(r.usuario.username) if r.usuario else ''}</td>
                </tr>
        """)

    html_content += ''.join(filas)
    html_content += """
            </tbody>
        </table>
    </body>
    </html>
    """

    pdf = HTML(string=html_content).write_pdf()

    response = make_response(pdf)
    response.headers['Content-Type'] = 'application/pdf'
    response.headers['Content-Disposition'] = 'attachment; filename=recepciones.pdf'
    return response


@recepcion_bp.route('/create', methods=['GET', 'POST'])
@login_required
@permiso_requerido('recepcion.crear', hacia='blueprint')
def create():

    form = RecepcionForm()
    form.tipoCert.choices = [(t.id, t.name) for t in TipoCertificado.query.filter_by(activo=True).order_by(TipoCertificado.name).all()]
    if form.validate_on_submit():
        try:
            total_folios = registrar_rango(
                fecha=form.fecha.data,
                anioCert=form.anioCert.data,
                tipoCert_id=form.tipoCert.data,
                folioInicial=form.folioInicial.data,
                folioFinal=form.folioFinal.data,
                user_id=current_user.id,
            )
            db.session.commit()
        except SolapeError as e:
            flash(str(e), 'danger')
            return redirect(url_for('recepcion.index'))
        except IntegrityError:
            db.session.rollback()
            flash('Conflicto de concurrencia al asignar el rango de recepción. '
                  + INTENTE_DE_NUEVO, 'danger')
            return redirect(url_for('recepcion.index'))

        flash(
            f'Recepción registrada. {total_folios} {plural(total_folios, "folio")} '
            f'{plural(total_folios, "creado")}.',
            'success',
        )
        return redirect(url_for('recepcion.index'))

    return render_template('recepcion/create.html', form=form)


def _validar_filas_recepcion(filas, usuario_default):
    """Valida cada fila del archivo contra la DB y entre sí.

    Devuelve (preview, payload): preview con una entrada por fila (etiqueta,
    responsable y error en español o None); payload solo con las filas válidas,
    saneado para el hidden de confirmación, con `usuario_id` por fila. La
    columna `usuario` manda: vacía o ausente → `usuario_default`; desconocida
    o inactiva → la fila queda inválida. Las filas también se cruzan entre sí
    porque dos rangos del mismo archivo pueden solaparse sin existir en la DB
    todavía.
    """
    tipos = {t.name.strip().lower(): t.id
             for t in TipoCertificado.query.filter_by(activo=True).all()}
    preview, payload, rangos = [], [], []
    for n, row in enumerate(filas, 1):
        fecha = import_excel.a_fecha(row.get('fecha'))
        anio = import_excel.a_entero(row.get('anioCert'))
        ini = import_excel.a_entero(row.get('folioInicial'))
        fin = import_excel.a_entero(row.get('folioFinal'))
        nombre = str(row.get('tipoCert') or '').strip()
        tipo_id = tipos.get(nombre.lower())

        error = None
        usuario = None
        if fecha is None:
            error = 'Fecha inválida.'
        elif anio is None or not 2000 <= anio <= 2100:
            error = 'Año fuera de rango (2000-2100).'
        elif tipo_id is None:
            error = f'Tipo de certificado desconocido: {nombre or "vacío"}.'
        elif ini is None or fin is None or ini < 1 or fin < 1:
            error = 'Número de folio inválido.'
        elif fin < ini:
            error = 'El folio final debe ser mayor o igual al inicial.'
        elif fin - ini + 1 > MAX_FOLIOS_POR_RECEPCION:
            error = f'El rango supera el máximo de {MAX_FOLIOS_POR_RECEPCION} folios.'
        if error is None:
            usuario, error = resolver_usuario_nombre(
                row.get(import_excel.COLUMNA_USUARIO), usuario_default)
        if error is None:
            solape = detectar_solape(anio, tipo_id, ini, fin)
            if solape:
                error = msg_solape(ini, fin, solape.folioInicial, solape.folioFinal)
        if error is None:
            for (a, t, i2, f2) in rangos:
                if a == anio and t == tipo_id and i2 <= fin and f2 >= ini:
                    error = 'El rango se solapa con otra fila del archivo.'
                    break
        if error is None:
            rangos.append((anio, tipo_id, ini, fin))
            payload.append({
                'fecha': fecha.isoformat(),
                'anioCert': anio,
                'tipoCert_id': tipo_id,
                'folioInicial': ini,
                'folioFinal': fin,
                'usuario_id': usuario.id,
            })
        preview.append({
            'n': n,
            'etiqueta': (
                f'{ini}-{fin} · {nombre or "?"} · {anio or "?"} · '
                f'{fecha.strftime("%d/%m/%Y") if fecha else "?"}'
            ),
            'usuario': usuario.username if usuario else None,
            'error': error,
        })
    return preview, payload


def _payload_recepcion(raw):
    """Hidden input del confirm → lista saneada, o None si hay datos ajenos."""
    try:
        datos = json.loads(raw or '')
    except (TypeError, ValueError):
        return None
    if not isinstance(datos, list) or not datos or len(datos) > MAX_RANGOS_POR_ARCHIVO:
        return None
    limpio = []
    for item in datos:
        if not isinstance(item, dict):
            return None
        try:
            fila = {
                'fecha': str(item['fecha']),
                'anioCert': int(item['anioCert']),
                'tipoCert_id': int(item['tipoCert_id']),
                'folioInicial': int(item['folioInicial']),
                'folioFinal': int(item['folioFinal']),
            }
        except (KeyError, TypeError, ValueError):
            return None
        # Opcional: payloads antiguos de sesiones abiertas no lo traen.
        uid = item.get('usuario_id')
        if uid is not None:
            try:
                uid = int(uid)
            except (TypeError, ValueError):
                return None
            if uid < 1:
                return None
        fila['usuario_id'] = uid
        total = fila['folioFinal'] - fila['folioInicial'] + 1
        if fila['folioInicial'] < 1 or not 2000 <= fila['anioCert'] <= 2100 \
                or total > MAX_FOLIOS_POR_RECEPCION:
            return None
        limpio.append(fila)
    return limpio


@recepcion_bp.route('/plantilla')
@login_required
@permiso_requerido('recepcion.importar')
def plantilla():
    tipo = TipoCertificado.query.filter_by(activo=True) \
        .order_by(TipoCertificado.name).first()
    ejemplo = [date.today(), 2026, tipo.name if tipo else 'Tipo de certificado',
               1, 100, current_user.username]
    return import_excel.respuesta_plantilla(
        'plantilla_recepcion.xlsx', HEADERS_RECEPCION, ejemplo)


@recepcion_bp.route('/importar', methods=['GET', 'POST'])
@login_required
@permiso_requerido('recepcion.importar')
def importar():
    form = ImportarRecepcionForm()
    if form.validate_on_submit():
        try:
            filas = import_excel.leer_filas(
                form.archivo.data, HEADERS_RECEPCION, MAX_RANGOS_POR_ARCHIVO,
                opcionales=(import_excel.COLUMNA_USUARIO,))
        except import_excel.ImportExcelError as e:
            flash(str(e), 'danger')
            return render_template('recepcion/importar.html', form=form)

        preview, payload = _validar_filas_recepcion(filas, current_user)
        total_validas = sum(1 for p in preview if p['error'] is None)
        if not total_validas:
            flash('Ninguna fila del archivo es válida.', 'danger')
        return render_template(
            'recepcion/importar.html', form=form, preview=preview,
            payload_json=json.dumps(payload), total_validas=total_validas,
            max_mostrar=MAX_PREVIA_MOSTRAR,
        )
    return render_template('recepcion/importar.html', form=form)


@recepcion_bp.route('/importar/confirmar', methods=['POST'])
@login_required
@permiso_requerido('recepcion.importar')
def importar_confirmar():
    payload = _payload_recepcion(request.form.get('payload'))
    if payload is None:
        flash('Datos de importación inválidos. Repita la operación.', 'danger')
        return redirect(url_for('recepcion.importar'))

    registrados = 0
    folios_creados = 0
    fallidos = []
    for item in payload:
        etiqueta = f"{item['folioInicial']}-{item['folioFinal']}"
        # uid inválido/inexistente → fallback al actor (payload es editable).
        usuario = usuario_activo(item.get('usuario_id'), current_user)
        try:
            total = registrar_rango(
                fecha=date.fromisoformat(item['fecha']),
                anioCert=item['anioCert'],
                tipoCert_id=item['tipoCert_id'],
                folioInicial=item['folioInicial'],
                folioFinal=item['folioFinal'],
                user_id=usuario.id,
                audit_user_id=current_user.id,
            )
            db.session.commit()
            registrados += 1
            folios_creados += total
        except (SolapeError, IntegrityError, ValueError, TypeError, KeyError) as e:
            db.session.rollback()
            if isinstance(e, SolapeError):
                fallidos.append(f'{etiqueta}: {e}')
            elif isinstance(e, IntegrityError):
                fallidos.append(f'{etiqueta}: conflicto de concurrencia.')
            else:
                fallidos.append(f'{etiqueta}: datos inválidos.')

    resumen = (f'{registrados} {plural(registrados, "recepción", "recepciones")} '
               f'{plural(registrados, "registrada")}, {folios_creados} '
               f'{plural(folios_creados, "folio")} {plural(folios_creados, "creado")}.')
    if fallidos:
        detalle = '; '.join(fallidos[:5]) + (' …' if len(fallidos) > 5 else '')
        flash(f'{resumen} {len(fallidos)} fila(s) fallaron: {detalle}', 'warning')
    elif not registrados:
        flash('Ninguna fila pudo importarse. Repita la operación.', 'danger')
    else:
        flash(resumen, 'success')
    return redirect(url_for('recepcion.index'))
