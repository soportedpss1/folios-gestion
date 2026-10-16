from flask import Blueprint, render_template, request, make_response, jsonify
from flask_login import login_required
from app.models.folio import Folio
from app.models.entrega_folio import EntregaFolio
from app.models.devolucion_folio import DevolucionFolio
from app.models.tipo_certificado import TipoCertificado
from app.models.centro import Centro
from sqlalchemy.orm import selectinload
from app.services.excel import neutralizar_formulas
from app.services.kpis import conteos_folios
from app.services.pdf import encabezado_pdf
from markupsafe import escape
from datetime import datetime
from io import BytesIO

reportes_bp = Blueprint('reportes', __name__, url_prefix='/reportes')


def _build_query(anio=None, tipo_id=None, centro_id=None, estado=None):
    query = Folio.query

    if anio:
        query = query.filter(Folio.anioCert == anio)

    if tipo_id:
        query = query.filter(Folio.tipoCert_id == tipo_id)

    if centro_id:
        query = query.join(EntregaFolio, Folio.entrega).filter(EntregaFolio.centroId == centro_id)

    if estado:
        if estado == 'nulo':
            query = query.filter(Folio.nulo == True)
        else:
            query = query.filter(Folio.estado == estado, Folio.nulo == False)

    return query


def _get_stats(anio=None, tipo_id=None, centro_id=None, estado=None):
    """Misma base que la tabla: si el filtro cambia, las tarjetas cambian."""
    return conteos_folios(_build_query(anio, tipo_id, centro_id, estado))


def _add_eager_loading(query):
    return query.options(
        selectinload(Folio.tipo_cert),
        selectinload(Folio.entrega)
            .selectinload(EntregaFolio.centro),
        selectinload(Folio.entrega)
            .selectinload(EntregaFolio.usuario),
        selectinload(Folio.devolucion)
            .selectinload(DevolucionFolio.centro),
        selectinload(Folio.devolucion)
            .selectinload(DevolucionFolio.usuario),
    )


@reportes_bp.route('/')
@login_required
def index():
    tipos = TipoCertificado.query.filter_by(activo=True).order_by(TipoCertificado.name).all()
    centros = Centro.query.filter_by(activo=True).order_by(Centro.name).all()

    # Año del filtro: ausente/vacío = "Todos" (el select no preselecciona);
    # el form es GET, así que al recargar debe respetar lo elegido.
    crudo = request.args.get('anioCert', '')
    try:
        current_anio = int(crudo) if crudo else None
    except ValueError:
        current_anio = None

    return render_template('reportes/index.html',
        folios=[],
        tipos=tipos,
        centros=centros,
        current_anio=current_anio,
        # El form es GET: al recargar hay que respetar lo elegido en TODOS los
        # selects (antes solo el año; tipo/centro/estado se perdían en "Todos").
        current_tipo=request.args.get('tipoCert', type=int),
        current_centro=request.args.get('centroId', type=int),
        current_estado=request.args.get('estado') or None
    )


@reportes_bp.route('/api/data')
@login_required
def api_data():
    anio = request.args.get('anioCert', type=int)
    tipo_id = request.args.get('tipoCert', type=int)
    centro_id = request.args.get('centroId', type=int)
    estado = request.args.get('estado')
    page = request.args.get('page', 1, type=int)
    per_page = min(max(request.args.get('per_page', 50, type=int), 1), 200)

    query = _build_query(anio, tipo_id, centro_id, estado)
    stats = _get_stats(anio, tipo_id, centro_id, estado)

    query = _add_eager_loading(query)
    pagination = query.order_by(Folio.folio, Folio.id).paginate(page=page, per_page=per_page, error_out=False)

    folios_data = []
    for f in pagination.items:
        folios_data.append({
            'folio': f.folio,
            'tipo_nombre': f.tipo_cert.name if f.tipo_cert else '',
            'tipo_color': f.tipo_cert.color if f.tipo_cert else '#0D6EFD',
            'anioCert': f.anioCert,
            'centro': f.entrega.centro.name if f.entrega and f.entrega.centro else None,
            'estado': f.estado,
            'digitado': f.digitado,
            'escaneado': f.escaneado,
            'nulo': f.nulo,
            'fechaEntrega': f.entrega.fechaEntrega.strftime('%d/%m/%Y') if f.entrega and f.entrega.fechaEntrega else None,
            'entregado_por': f.entrega.usuario.username if f.entrega and f.entrega.usuario else None,
            'fechaDevolucion': f.devolucion.fechaDevolucion.strftime('%d/%m/%Y') if f.devolucion and f.devolucion.fechaDevolucion else None,
            'recibido_por': f.devolucion.usuario.username if f.devolucion and f.devolucion.usuario else None,
        })

    return jsonify({
        'folios': folios_data,
        'stats': stats,
        'pagination': {
            'page': pagination.page,
            'pages': pagination.pages,
            'total': pagination.total,
            'has_next': pagination.has_next,
            'has_prev': pagination.has_prev,
            'next_num': pagination.next_num,
            'prev_num': pagination.prev_num,
        }
    })


# Columnas compartidas entre el PDF y el Excel del reporte: el mismo
# reporte no puede cambiar de orden ni de contenido según el formato
# (auditoría UI/UX: 12 columnas en Excel, 11 y otro orden en el PDF).
COLUMNAS_REPORTE = [
    'Folio', 'Tipo Certificado', 'Año', 'Centro', 'Estado',
    'Digitado', 'Escaneado', 'Nulo',
    'Fecha Entrega', 'Entregado por',
    'Fecha Devolución', 'Recibido Por',
]


def fila_reporte(f):
    """Valores de una fila en orden COLUMNAS_REPORTE; idénticos en PDF y Excel."""
    entrega = f.entrega
    devolucion = f.devolucion
    return [
        f.folio,
        f.tipo_cert.name if f.tipo_cert else '',
        f.anioCert,
        entrega.centro.name if entrega and entrega.centro else '',
        f.estado,
        'Sí' if f.digitado else 'No',
        'Sí' if f.escaneado else 'No',
        'Sí' if f.nulo else 'No',
        entrega.fechaEntrega.strftime('%d/%m/%Y')
        if entrega and entrega.fechaEntrega else '-',
        entrega.usuario.username if entrega and entrega.usuario else '-',
        devolucion.fechaDevolucion.strftime('%d/%m/%Y')
        if devolucion and devolucion.fechaDevolucion else '-',
        devolucion.usuario.username if devolucion and devolucion.usuario else '-',
    ]


@reportes_bp.route('/export/excel')
@login_required
def export_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    anio = request.args.get('anioCert', type=int)
    tipo_id = request.args.get('tipoCert', type=int)
    centro_id = request.args.get('centroId', type=int)
    estado = request.args.get('estado')

    query = _build_query(anio, tipo_id, centro_id, estado)
    query = _add_eager_loading(query)
    folios = query.order_by(Folio.folio, Folio.id).all()

    wb = Workbook()
    ws = wb.active
    ws.title = f'Folios {anio or "Todos"}'

    headers = COLUMNAS_REPORTE

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
    header_align = Alignment(horizontal='center')

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for row, f in enumerate(folios, 2):
        for col, valor in enumerate(fila_reporte(f), 1):
            ws.cell(row=row, column=col, value=valor)

    for col in ws.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2

    output = BytesIO()
    neutralizar_formulas(wb)  # nunca escribir '=...' como fórmula (inyección al abrir)
    wb.save(output)
    output.seek(0)

    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f'attachment; filename=folios_{anio or "todos"}.xlsx'
    return response


@reportes_bp.route('/export/pdf')
@login_required
def export_pdf():
    from weasyprint import HTML

    anio = request.args.get('anioCert', type=int)
    tipo_id = request.args.get('tipoCert', type=int)
    centro_id = request.args.get('centroId', type=int)
    estado = request.args.get('estado')

    query = _build_query(anio, tipo_id, centro_id, estado)
    query = _add_eager_loading(query)
    folios = query.order_by(Folio.folio, Folio.id).all()

    encabezado = encabezado_pdf(f'Reporte de Folios - Año {anio or "Todos"}')
    thead = ''.join(f'<th>{columna}</th>' for columna in COLUMNAS_REPORTE)
    html_content = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <style>
            @page {{ size: landscape; margin: 1.5cm; }}
            body {{ font-family: Arial, sans-serif; margin: 20px; font-size: 10px; }}
            h1 {{ color: #0d6efd; text-align: center; font-size: 16px; }}
            h2 {{ color: #333; text-align: center; font-size: 13px; margin-top: 5px; }}
            h3 {{ color: #666; text-align: center; font-size: 11px; margin-top: 5px; }}
            table {{ width: 100%; border-collapse: collapse; margin-top: 15px; font-size: 9px; }}
            th, td {{ border: 1px solid #ddd; padding: 5px 6px; text-align: left; }}
            th {{ background-color: #0d6efd; color: white; font-size: 9px; }}
            tr:nth-child(even) {{ background-color: #f2f2f2; }}
            .header {{ text-align: center; margin-bottom: 20px; }}
            .date {{ color: #666; font-size: 10px; text-align: right; margin-bottom: 10px; }}
        </style>
    </head>
    <body>
        {encabezado}
        <p class="date">Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}</p>
        <table>
            <thead>
                <tr>{thead}</tr>
            </thead>
            <tbody>
    """

    filas = []
    for f in folios:
        filas.append('<tr>' + ''.join(
            f'<td>{escape(str(valor))}</td>' for valor in fila_reporte(f))
            + '</tr>')

    html_content += ''.join(filas)
    html_content += """
            </tbody>
        </table>
    </body>
    </html>
    """

    pdf = HTML(string=html_content).write_pdf()

    response = make_response(pdf)
    response.headers['Content-Type'] = 'application/pdf'
    response.headers['Content-Disposition'] = f'attachment; filename=folios_{anio or "todos"}.pdf'
    return response
