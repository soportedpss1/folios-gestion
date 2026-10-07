import json

from flask import Blueprint, render_template, redirect, url_for, flash, make_response, request
from flask_login import login_required, current_user
from app.extensions import db
from app.models.centro import Centro
from app.models.folio import Folio
from app.models.entrega_folio import EntregaFolio
from app.models.devolucion_folio import DevolucionFolio
from app.centros.forms import CentroForm, ImportarCentroForm
from app.services import import_excel
from app.services.audit import log_audit
from app.services.pdf import encabezado_pdf, logo_pdf
from app.services.queries import get_or_404
from app.services.excel import neutralizar_formulas
from app.decorators import permiso_requerido
from app.utils import plural
from markupsafe import escape
from datetime import datetime, date
from io import BytesIO
from sqlalchemy.orm import selectinload


centros_bp = Blueprint('centros', __name__, url_prefix='/centros')

# Importación de centros: columnas exactas de la plantilla y tope de filas.
HEADERS_CENTROS = ['nombre', 'telefono', 'direccion']
MAX_CENTROS_POR_ARCHIVO = 500

# Carga anticipada de lo que los listados/exportaciones recorren fila a fila.
# Se construyen en runtime: los backrefs (folio/centro/usuario) se configuran
# al primer uso de los mappers, no al importar el módulo.
def _eager_entrega():
    return (
        selectinload(EntregaFolio.folio).selectinload(Folio.tipo_cert),
        selectinload(EntregaFolio.folio).selectinload(Folio.devolucion)
            .selectinload(DevolucionFolio.usuario),
        selectinload(EntregaFolio.usuario),
    )


def _eager_devolucion():
    return (
        selectinload(DevolucionFolio.folio).selectinload(Folio.tipo_cert),
        selectinload(DevolucionFolio.folio).selectinload(Folio.entrega),
        selectinload(DevolucionFolio.usuario),
    )


def _datos_del_centro(centro_id):
    """Entregas, pendientes y devoluciones de un centro: detail y ambos exports."""
    entregas = EntregaFolio.query.options(*_eager_entrega())\
        .filter_by(centroId=centro_id).order_by(EntregaFolio.fechaEntrega.desc()).all()

    folios_devueltos_ids = set(d.folio_id for d in
        DevolucionFolio.query.filter_by(centroId=centro_id).all())

    pendientes = [e for e in entregas if e.folio_id not in folios_devueltos_ids]

    devueltos = DevolucionFolio.query.options(*_eager_devolucion())\
        .filter_by(centroId=centro_id).order_by(DevolucionFolio.fechaDevolucion.desc()).all()

    return entregas, pendientes, devueltos


@centros_bp.route('/')
@login_required
def index():
    centros = Centro.query.order_by(Centro.name).all()
    return render_template('centros/index.html', centros=centros)


@centros_bp.route('/<int:id>/detail')
@login_required
def detail(id):
    centro = get_or_404(Centro, id)

    entregas, pendientes, devueltos = _datos_del_centro(id)

    hoy = date.today()

    return render_template('centros/detail.html',
        centro=centro,
        entregas=entregas,
        pendientes=pendientes,
        devueltos=devueltos,
        hoy=hoy
    )


@centros_bp.route('/create', methods=['GET', 'POST'])
@permiso_requerido('centros.gestionar')
def create():
    form = CentroForm()
    if form.validate_on_submit():
        centro = Centro(
            name=form.name.data,
            telefono=form.telefono.data,
            direccion=form.direccion.data
        )
        db.session.add(centro)
        db.session.flush()
        log_audit(current_user.id, 'INSERT', 'centros', centro.id, None, centro.to_dict())
        db.session.commit()
        flash('Centro creado exitosamente.', 'success')
        return redirect(url_for('centros.index'))
    return render_template('centros/create.html', form=form)


@centros_bp.route('/<int:id>/edit', methods=['GET', 'POST'])
@permiso_requerido('centros.gestionar')
def edit(id):
    centro = get_or_404(Centro, id)
    datos_anteriores = centro.to_dict()
    form = CentroForm(obj=centro)
    if form.validate_on_submit():
        centro.name = form.name.data
        centro.telefono = form.telefono.data
        centro.direccion = form.direccion.data
        log_audit(current_user.id, 'UPDATE', 'centros', centro.id, datos_anteriores, centro.to_dict())
        db.session.commit()
        flash('Centro actualizado exitosamente.', 'success')
        return redirect(url_for('centros.index'))
    return render_template('centros/edit.html', form=form, centro=centro)


@centros_bp.route('/<int:id>/delete', methods=['POST'])
@permiso_requerido('centros.gestionar')
def delete(id):
    centro = get_or_404(Centro, id)
    datos_anteriores = centro.to_dict()
    centro.activo = False
    log_audit(current_user.id, 'DELETE', 'centros', centro.id, datos_anteriores, None)
    db.session.commit()
    flash('Centro desactivado exitosamente.', 'success')
    return redirect(url_for('centros.index'))


def _nombres_existentes():
    return {nombre.strip().lower()
            for (nombre,) in db.session.query(Centro.name).all() if nombre}


def _validar_filas_centros(filas):
    """Valida cada fila contra la DB y entre sí.

    Devuelve (preview, payload): preview con una entrada por fila; payload solo
    con las filas válidas, saneado para el hidden de confirmación. El nombre se
    compara en minúsculas sin distinguir mayúsculas: duplicado en la DB (incluye
    inactivos) o repetido dentro del archivo → fila inválida.
    """
    existentes = _nombres_existentes()
    vistos = set()
    preview, payload = [], []
    for n, row in enumerate(filas, 1):
        nombre = str(row.get('nombre') or '').strip()
        telefono = str(row.get('telefono') or '').strip()
        direccion = str(row.get('direccion') or '').strip()
        clave = nombre.lower()

        error = None
        if not nombre:
            error = 'Nombre vacío.'
        elif not 3 <= len(nombre) <= 255:
            error = 'El nombre debe tener entre 3 y 255 caracteres.'
        elif len(telefono) > 50:
            error = 'El teléfono no puede superar 50 caracteres.'
        elif len(direccion) > 255:
            error = 'La dirección no puede superar 255 caracteres.'
        elif clave in existentes:
            error = f'Centro ya existe: {nombre}.'
        elif clave in vistos:
            error = 'Nombre repetido en el archivo.'

        if error is None:
            vistos.add(clave)
            payload.append({'nombre': nombre, 'telefono': telefono,
                            'direccion': direccion})
        preview.append({'n': n, 'nombre': nombre or '?',
                        'telefono': telefono, 'direccion': direccion,
                        'error': error})
    return preview, payload


def _payload_centros(raw):
    """Hidden input del confirm → lista saneada, o None si hay datos ajenos."""
    try:
        datos = json.loads(raw or '')
    except (TypeError, ValueError):
        return None
    if not isinstance(datos, list) or not datos or len(datos) > MAX_CENTROS_POR_ARCHIVO:
        return None
    limpio = []
    for item in datos:
        if not isinstance(item, dict):
            return None
        try:
            nombre = str(item['nombre']).strip()
            telefono = str(item['telefono']).strip()
            direccion = str(item['direccion']).strip()
        except (KeyError, TypeError):
            return None
        if not 3 <= len(nombre) <= 255 or len(telefono) > 50 or len(direccion) > 255:
            return None
        limpio.append({'nombre': nombre, 'telefono': telefono,
                       'direccion': direccion})
    return limpio


@centros_bp.route('/plantilla')
@login_required
@permiso_requerido('centros.importar')
def plantilla():
    ejemplo = ['Hospital General', '555555', 'Calle 1 #2']
    return import_excel.respuesta_plantilla(
        'plantilla_centros.xlsx', HEADERS_CENTROS, ejemplo)


@centros_bp.route('/importar', methods=['GET', 'POST'])
@login_required
@permiso_requerido('centros.importar')
def importar():
    form = ImportarCentroForm()
    if form.validate_on_submit():
        try:
            filas = import_excel.leer_filas(
                form.archivo.data, HEADERS_CENTROS, MAX_CENTROS_POR_ARCHIVO)
        except import_excel.ImportExcelError as e:
            flash(str(e), 'danger')
            return render_template('centros/importar.html', form=form)

        preview, payload = _validar_filas_centros(filas)
        total_validas = sum(1 for p in preview if p['error'] is None)
        if not total_validas:
            flash('Ninguna fila del archivo es válida.', 'danger')
        return render_template(
            'centros/importar.html', form=form, preview=preview,
            payload_json=json.dumps(payload), total_validas=total_validas,
            max_mostrar=200,
        )
    return render_template('centros/importar.html', form=form)


@centros_bp.route('/importar/confirmar', methods=['POST'])
@login_required
@permiso_requerido('centros.importar')
def importar_confirmar():
    payload = _payload_centros(request.form.get('payload'))
    if payload is None:
        flash('Datos de importación inválidos. Repita la operación.', 'danger')
        return redirect(url_for('centros.importar'))

    # Revalida contra la DB: el hidden es editable por el usuario.
    existentes = _nombres_existentes()
    creados = 0
    fallidos = []
    for item in payload:
        clave = item['nombre'].lower()
        if clave in existentes:
            fallidos.append(f"{item['nombre']}: ya existe.")
            continue
        try:
            centro = Centro(name=item['nombre'], telefono=item['telefono'],
                            direccion=item['direccion'], activo=True)
            db.session.add(centro)
            db.session.flush()
            log_audit(current_user.id, 'INSERT', 'centros', centro.id, None,
                      centro.to_dict())
            db.session.commit()
            existentes.add(clave)
            creados += 1
        except Exception:
            db.session.rollback()
            fallidos.append(f"{item['nombre']}: conflicto al guardar.")

    resumen = (f'{creados} {plural(creados, "centro")} '
               f'{plural(creados, "creado")}.')
    if fallidos:
        detalle = '; '.join(fallidos[:5]) + (' …' if len(fallidos) > 5 else '')
        flash(f'{resumen} {len(fallidos)} fila(s) fallaron: {detalle}', 'warning')
    elif not creados:
        flash('Ningún centro pudo importarse. Repita la operación.', 'danger')
    else:
        flash(resumen, 'success')
    return redirect(url_for('centros.index'))


@centros_bp.route('/export/excel')
@login_required
def export_lista_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    centros = Centro.query.order_by(Centro.name).all()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Centros'

    headers = ['#', 'Nombre', 'Teléfono', 'Dirección', 'Estado']
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
    header_align = Alignment(horizontal='center')
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for row, c in enumerate(centros, 2):
        ws.cell(row=row, column=1, value=c.id)
        ws.cell(row=row, column=2, value=c.name)
        ws.cell(row=row, column=3, value=c.telefono or '')
        ws.cell(row=row, column=4, value=c.direccion or '')
        ws.cell(row=row, column=5, value='Activo' if c.activo else 'Inactivo')

    for col in ws.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2

    output = BytesIO()
    neutralizar_formulas(wb)  # nunca escribir '=...' como fórmula (inyección al abrir)
    wb.save(output)
    output.seek(0)

    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = 'attachment; filename=centros.xlsx'
    return response


@centros_bp.route('/export/pdf')
@login_required
def export_lista_pdf():
    from weasyprint import HTML

    centros = Centro.query.order_by(Centro.name).all()

    encabezado = encabezado_pdf('Gestion de Centros de Salud')
    html_content = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <style>
            @page {{ size: landscape; margin: 1.5cm; }}
            body {{ font-family: Arial, sans-serif; margin: 20px; font-size: 10px; }}
            h1 {{ color: #0d6efd; text-align: center; font-size: 16px; }}
            h2 {{ color: #333; text-align: center; font-size: 13px; margin-top: 5px; }}
            h3 {{ color: #666; text-align: center; font-size: 11px; margin-top: 5px; }}
            table {{ width: 100%; border-collapse: collapse; margin-top: 15px; font-size: 10px; }}
            th, td {{ border: 1px solid #ddd; padding: 5px 6px; text-align: left; }}
            th {{ background-color: #0d6efd; color: white; font-size: 10px; }}
            tr:nth-child(even) {{ background-color: #f2f2f2; }}
            .header {{ text-align: center; margin-bottom: 20px; }}
            .date {{ color: #666; font-size: 10px; text-align: right; margin-bottom: 10px; }}
        </style>
    </head>
    <body>
        {encabezado}
        <p class="date">Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}</p>
        <table>
            <thead>
                <tr>
                    <th>#</th>
                    <th>Nombre</th>
                    <th>Telefono</th>
                    <th>Direccion</th>
                    <th>Estado</th>
                </tr>
            </thead>
            <tbody>
    """

    filas = []
    for c in centros:
        filas.append(f"""
                <tr>
                    <td>{c.id}</td>
                    <td>{escape(c.name)}</td>
                    <td>{escape(c.telefono or '-')}</td>
                    <td>{escape(c.direccion or '-')}</td>
                    <td>{'Activo' if c.activo else 'Inactivo'}</td>
                </tr>
        """)

    html_content += ''.join(filas)
    html_content += """
            </tbody>
        </table>
    </body>
    </html>
    """

    pdf = HTML(string=html_content).write_pdf()

    response = make_response(pdf)
    response.headers['Content-Type'] = 'application/pdf'
    response.headers['Content-Disposition'] = 'attachment; filename=centros.pdf'
    return response


@centros_bp.route('/<int:id>/export/excel')
@login_required
def export_excel(id):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    centro = get_or_404(Centro, id)

    entregas, pendientes, devueltos = _datos_del_centro(id)

    hoy = date.today()

    wb = Workbook()

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
    header_align = Alignment(horizontal='center')

    # Sheet 1: Pendientes
    ws1 = wb.active
    ws1.title = 'Pendientes'
    headers1 = ['Folio', 'Tipo Certificado', 'Fecha Entrega', 'Entregado por', 'Días en Centro']
    for col, header in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    for row, e in enumerate(pendientes, 2):
        dias = (hoy - e.fechaEntrega).days
        ws1.cell(row=row, column=1, value=e.folio.folio)
        ws1.cell(row=row, column=2, value=e.folio.tipo_cert.name if e.folio.tipo_cert else '')
        ws1.cell(row=row, column=3, value=e.fechaEntrega.strftime('%d/%m/%Y'))
        ws1.cell(row=row, column=4, value=e.usuario.username if e.usuario else '')
        ws1.cell(row=row, column=5, value=dias)
    for col in ws1.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        ws1.column_dimensions[col[0].column_letter].width = max_length + 2

    # Sheet 2: Todos Entregados
    ws2 = wb.create_sheet('Todos Entregados')
    headers2 = ['Folio', 'Tipo Certificado', 'Estado', 'F. Entrega', 'Entregado por', 'F. Devolución', 'Recibido Por']
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    for row, e in enumerate(entregas, 2):
        folio = e.folio
        ws2.cell(row=row, column=1, value=folio.folio)
        ws2.cell(row=row, column=2, value=folio.tipo_cert.name if folio.tipo_cert else '')
        ws2.cell(row=row, column=3, value=folio.estado)
        ws2.cell(row=row, column=4, value=e.fechaEntrega.strftime('%d/%m/%Y'))
        ws2.cell(row=row, column=5, value=e.usuario.username if e.usuario else '')
        if folio.devolucion:
            ws2.cell(row=row, column=6, value=folio.devolucion.fechaDevolucion.strftime('%d/%m/%Y') if folio.devolucion.fechaDevolucion else '-')
            ws2.cell(row=row, column=7, value=folio.devolucion.usuario.username if folio.devolucion.usuario else '-')
        else:
            ws2.cell(row=row, column=6, value='-')
            ws2.cell(row=row, column=7, value='-')
    for col in ws2.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        ws2.column_dimensions[col[0].column_letter].width = max_length + 2

    # Sheet 3: Devueltos
    ws3 = wb.create_sheet('Devueltos')
    headers3 = ['Folio', 'Tipo Certificado', 'F. Entrega', 'F. Devolución', 'Recibido Por', 'Días en Centro']
    for col, header in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    for row, d in enumerate(devueltos, 2):
        dias = (d.fechaDevolucion - d.folio.entrega.fechaEntrega).days if d.folio.entrega else 0
        ws3.cell(row=row, column=1, value=d.folio.folio)
        ws3.cell(row=row, column=2, value=d.folio.tipo_cert.name if d.folio.tipo_cert else '')
        ws3.cell(row=row, column=3, value=d.folio.entrega.fechaEntrega.strftime('%d/%m/%Y') if d.folio.entrega else '-')
        ws3.cell(row=row, column=4, value=d.fechaDevolucion.strftime('%d/%m/%Y'))
        ws3.cell(row=row, column=5, value=d.usuario.username if d.usuario else '')
        ws3.cell(row=row, column=6, value=dias)
    for col in ws3.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        ws3.column_dimensions[col[0].column_letter].width = max_length + 2

    output = BytesIO()
    neutralizar_formulas(wb)  # nunca escribir '=...' como fórmula (inyección al abrir)
    wb.save(output)
    output.seek(0)

    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f'attachment; filename=centro_{centro.name}_{date.today().strftime("%Y%m%d")}.xlsx'
    return response


@centros_bp.route('/<int:id>/export/pdf')
@login_required
def export_pdf(id):
    from weasyprint import HTML

    centro = get_or_404(Centro, id)

    entregas, pendientes, devueltos = _datos_del_centro(id)

    hoy = date.today()

    def build_pendientes_rows():
        filas = []
        for e in pendientes:
            dias = (hoy - e.fechaEntrega).days
            filas.append(f"""
                <tr>
                    <td>{e.folio.folio}</td>
                    <td>{escape(e.folio.tipo_cert.name) if e.folio.tipo_cert else ''}</td>
                    <td>{e.fechaEntrega.strftime('%d/%m/%Y')}</td>
                    <td>{escape(e.usuario.username) if e.usuario else ''}</td>
                    <td>{dias} días</td>
                </tr>
            """)
        if not filas:
            return '<tr><td colspan="5" style="text-align:center">No hay folios pendientes</td></tr>'
        return ''.join(filas)

    def build_entregas_rows():
        filas = []
        for e in entregas:
            folio = e.folio
            fecha_dev = folio.devolucion.fechaDevolucion.strftime('%d/%m/%Y') if folio.devolucion and folio.devolucion.fechaDevolucion else '-'
            recv_por = folio.devolucion.usuario.username if folio.devolucion and folio.devolucion.usuario else '-'
            filas.append(f"""
                <tr>
                    <td>{folio.folio}</td>
                    <td>{escape(folio.tipo_cert.name) if folio.tipo_cert else ''}</td>
                    <td>{escape(folio.estado)}</td>
                    <td>{e.fechaEntrega.strftime('%d/%m/%Y')}</td>
                    <td>{escape(e.usuario.username) if e.usuario else ''}</td>
                    <td>{fecha_dev}</td>
                    <td>{escape(recv_por)}</td>
                </tr>
            """)
        if not filas:
            return '<tr><td colspan="7" style="text-align:center">No hay entregas registradas</td></tr>'
        return ''.join(filas)

    def build_devueltos_rows():
        filas = []
        for d in devueltos:
            dias = (d.fechaDevolucion - d.folio.entrega.fechaEntrega).days if d.folio.entrega else 0
            fecha_ent = d.folio.entrega.fechaEntrega.strftime('%d/%m/%Y') if d.folio.entrega else '-'
            filas.append(f"""
                <tr>
                    <td>{d.folio.folio}</td>
                    <td>{escape(d.folio.tipo_cert.name) if d.folio.tipo_cert else ''}</td>
                    <td>{fecha_ent}</td>
                    <td>{d.fechaDevolucion.strftime('%d/%m/%Y')}</td>
                    <td>{escape(d.usuario.username) if d.usuario else ''}</td>
                    <td>{dias} días</td>
                </tr>
            """)
        if not filas:
            return '<tr><td colspan="6" style="text-align:center">No hay devoluciones registradas</td></tr>'
        return ''.join(filas)

    logo = logo_pdf()
    html_content = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <style>
            @page {{ size: landscape; margin: 1.5cm; }}
            body {{ font-family: Arial, sans-serif; margin: 20px; font-size: 10px; }}
            h1 {{ color: #0d6efd; text-align: center; font-size: 16px; }}
            h2 {{ color: #333; text-align: center; font-size: 13px; margin-top: 5px; }}
            h3 {{ color: #666; text-align: center; font-size: 11px; margin-top: 5px; }}
            table {{ width: 100%; border-collapse: collapse; margin-top: 10px; margin-bottom: 20px; font-size: 9px; }}
            th, td {{ border: 1px solid #ddd; padding: 5px 6px; text-align: left; }}
            th {{ background-color: #0d6efd; color: white; font-size: 9px; }}
            tr:nth-child(even) {{ background-color: #f2f2f2; }}
            .section-title {{ font-size: 12px; font-weight: bold; color: #0d6efd; margin-top: 15px; margin-bottom: 5px; }}
            .info-table td {{ border: none; padding: 2px 10px; }}
            .info-table tr td:first-child {{ font-weight: bold; color: #666; }}
        </style>
    </head>
    <body>
        {logo}
        <h1>Direccion Provincial de Salud Santiago 1</h1>
        <h2>Departamento de Bioestadistica</h2>
        <h3>Detalle del Centro: {escape(centro.name)}</h3>
        <p style="color:#666; font-size:10px; text-align:right;">Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}</p>

        <table class="info-table">
            <tr><td>Nombre:</td><td>{escape(centro.name)}</td></tr>
            <tr><td>Telefono:</td><td>{escape(centro.telefono) if centro.telefono else '-'}</td></tr>
            <tr><td>Direccion:</td><td>{escape(centro.direccion) if centro.direccion else '-'}</td></tr>
        </table>

        <div class="section-title">Pendientes de Devolucion ({len(pendientes)})</div>
        <table>
            <thead>
                <tr>
                    <th>Folio</th>
                    <th>Tipo Certificado</th>
                    <th>Fecha Entrega</th>
                    <th>Entregado por</th>
                    <th>Dias en Centro</th>
                </tr>
            </thead>
            <tbody>
                {build_pendientes_rows()}
            </tbody>
        </table>

        <div class="section-title">Todos Entregados ({len(entregas)})</div>
        <table>
            <thead>
                <tr>
                    <th>Folio</th>
                    <th>Tipo Certificado</th>
                    <th>Estado</th>
                    <th>F. Entrega</th>
                    <th>Entregado por</th>
                    <th>F. Devolucion</th>
                    <th>Recibido Por</th>
                </tr>
            </thead>
            <tbody>
                {build_entregas_rows()}
            </tbody>
        </table>

        <div class="section-title">Devueltos ({len(devueltos)})</div>
        <table>
            <thead>
                <tr>
                    <th>Folio</th>
                    <th>Tipo Certificado</th>
                    <th>F. Entrega</th>
                    <th>F. Devolucion</th>
                    <th>Recibido Por</th>
                    <th>Dias en Centro</th>
                </tr>
            </thead>
            <tbody>
                {build_devueltos_rows()}
            </tbody>
        </table>
    </body>
    </html>
    """

    pdf = HTML(string=html_content).write_pdf()

    response = make_response(pdf)
    response.headers['Content-Type'] = 'application/pdf'
    response.headers['Content-Disposition'] = f'attachment; filename=centro_{centro.name}_{date.today().strftime("%Y%m%d")}.pdf'
    return response
